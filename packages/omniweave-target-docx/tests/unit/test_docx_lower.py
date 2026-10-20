"""`compile.docx.native`'s `lower()` and `postflight()`: owflow/1 to a `.docx`. D782.

Every test builds a JobDir in `tmp_path` (`job.json`, `il/`, `assets/index.json`), lowers it with
`DocxCompiler` over a `FrozenLedger`, and reads the package back with `zipfile`, lxml and
python-docx. The determinism tests compile twice in one process and once in a child with another
`PYTHONHASHSEED`, and compare part digests.
"""

from __future__ import annotations

import hashlib
import io
import json
import os
import struct
import subprocess  # noqa: TID251 - a second interpreter is the only witness for PYTHONHASHSEED.
import sys
import tomllib
import zipfile
import zlib
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any

import docx
import pytest
from docx.image.image import Image
from lxml import etree  # type: ignore[attr-defined] -- lxml ships no stubs
from omniweave_conform.harness import MemoryBlobStore, make_io
from omniweave_core.errors import TargetError
from omniweave_core.model.block import Cite
from omniweave_core.out.brief import BRIEF_SCHEMA, parse_brief
from omniweave_core.out.compile import lower_job
from omniweave_core.out.constructs import ConstructTable
from omniweave_core.out.gate import GateStatus
from omniweave_core.out.job import ILBundle, Job, JobDeterminism
from omniweave_core.out.ledger import FrozenLedger, LedgerRow, ledger_digest, slot, write_index
from omniweave_core.out.opc import part_digest, write_opc
from omniweave_core.out.plan import Plan, UnitOutcome
from omniweave_core.out.receipt import Receipt
from omniweave_ports import ArtifactKind
from omniweave_target_docx.driver import HERE, NORMALIZERS, PROVENANCE, DocxCompiler, table
from omniweave_target_docx.lower import _pack
from omniweave_target_docx.lower._emit import resolve

TABLE = table()
CLOCK_NS = 1_700_000_000_123_456_789
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
TRACEPARENT = "00-" + "1" * 32 + "-" + "2" * 16 + "-01"


def png(width: int, height: int) -> bytes:
    raw = b"".join(b"\x00" + b"\x10\x20\x30" * width for _ in range(height))

    def chunk(kind: bytes, data: bytes) -> bytes:
        crc = struct.pack(">I", zlib.crc32(kind + data))
        return struct.pack(">I", len(data)) + kind + data + crc

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


WIDE = png(2000, 100)
SMALL = png(40, 20)


def _brief(canvas: str) -> Any:
    return parse_brief(
        f"{BRIEF_SCHEMA}\n## canvas\n- format: {canvas}\n- unit: section\n## communication\n"
        "- primary_language: en-GB\n## direction\n## density\n- per_unit: light\n## motion\n"
        "- level: static\n## typography\n## rhythm\n## forbidden\n",
        baseline=(),
    )


@dataclass(frozen=True, slots=True)
class Lowered:
    data: bytes
    members: Mapping[str, bytes]
    plan: Plan
    path: Path

    def xml(self, name: str = "word/document.xml") -> Any:
        return etree.fromstring(self.members[name])

    def document(self) -> Any:
        return docx.Document(io.BytesIO(self.data))

    def receipt(self) -> Receipt:
        return DocxCompiler().postflight(str(self.path), self.plan)


def _job(
    tmp_path: Path,
    units: Sequence[str],
    *,
    assets: Sequence[bytes] = (),
    canvas: str = "A4",
    plan_units: Sequence[str] | None = None,
    cites: Mapping[str, int] | None = None,
) -> tuple[Path, Plan, FrozenLedger, MemoryBlobStore]:
    job = tmp_path / "job"
    il = job / "il"
    il.mkdir(parents=True)
    rows = "".join(
        f'\n[[unit]]\nid = "u{n:02d}"\nord = {n}\nfile = "u{n:02d}.html"\n'
        for n in range(1, len(units) + 1)
    )
    (il / "il.toml").write_text(
        f'schema = "omniweave.out-il/1"\nil = "owflow/1"\nconstructs_digest = "{TABLE.digest()}"\n'
        + rows,
        encoding="utf-8",
    )
    for n, text in enumerate(units, 1):
        (il / f"u{n:02d}.html").write_text(text, encoding="utf-8", newline="")
    store = MemoryBlobStore(tmp_path / "cas")
    frozen = [
        LedgerRow(n, (sha := store.add(blob)), slot(sha), "image", "u01")
        for n, blob in enumerate(assets, 1)
    ]
    write_index(job, frozen)
    compiler = DocxCompiler()
    bundle = ILBundle.load(il, carrier=compiler.carrier, constructs=TABLE)
    brief = _brief(canvas)
    roster = bundle.units
    if plan_units is not None:
        roster = tuple(u for u in bundle.units if u.id in plan_units)
    plan = Plan(
        schema="omniweave.out-plan/1",
        target="docx",
        slug="memo",
        artifact_gen=1,
        delivery="draft",
        units=roster,
        unit_count=len(roster),
        brief=brief,
        corpus="acme",
        cites={Cite(k): v for k, v in (cites or {}).items()},
        gate_status=GateStatus.MANUAL_REQUIRED,
        il_digest=bundle.il_digest,
        constructs_digest=TABLE.digest(),
        ledger_digest=ledger_digest(frozen),
        brief_digest=brief.digest(),
        card_sha256="sha256:" + "4" * 64,
    )
    Job("omniweave.out-job/1", plan, TRACEPARENT, JobDeterminism(CLOCK_NS)).write(job)
    return job, plan, FrozenLedger(frozen, store), store


def lowered(tmp_path: Path, *units: str, **kwargs: Any) -> Lowered:
    job, plan, ledger, store = _job(tmp_path, units, **kwargs)
    compiler = DocxCompiler()
    bundle = ILBundle.load(job / "il", carrier=compiler.carrier, constructs=TABLE)
    result = compiler.lower(plan, bundle, ledger, make_io(store, tmp_path / "io"))
    [ref] = result.produced
    assert ref.kind == ArtifactKind.ARTIFACT
    data = ref.inline if ref.inline is not None else store.open(ref.blob or "").read()
    path = tmp_path / "memo.docx"
    path.write_bytes(data)
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        members = {info.filename: zf.read(info) for info in zf.infolist()}
    return Lowered(data, members, plan, path)


def section(body: str, el: str = "") -> str:
    address = f' data-ow-el="{el}"' if el else ""
    return f"<section{address}>{body}</section>"


def texts(lowered_: Lowered) -> list[str]:
    return [p.text for p in lowered_.document().paragraphs if p.text]


def _ptext(p: Any) -> str:
    return "".join(t.text or "" for t in p.iter(f"{W}t"))


# ---------------------------------------------------------------------------------------------
# OUT1: the table's maps_to and lower/'s public functions, both ways
# ---------------------------------------------------------------------------------------------


def test_every_maps_to_resolves() -> None:
    """D783 lowered the last four: notes and the provenance part."""
    named = {row.maps_to for row in TABLE.rows if row.maps_to}
    assert {m for m in named if resolve(m, _pack.FAMILIES) is None} == set()


def test_every_public_lowering_function_is_one_a_row_names() -> None:
    """OUT1's other direction: `lower/` has no function the table does not declare."""
    named = {row.maps_to for row in TABLE.rows if row.maps_to}
    for family, module in _pack.FAMILIES.items():
        exported = {f"{family}:{n}" for n in module.__all__ if callable(getattr(module, n))}
        assert exported <= named, (family, sorted(exported - named))
        assert exported == {m for m in named if m.startswith(f"{family}:")}


# ---------------------------------------------------------------------------------------------
# preflight (09 section 5.1)
# ---------------------------------------------------------------------------------------------


def test_a_row_whose_function_is_not_built_is_refused_before_a_byte(tmp_path: Path) -> None:
    """Preflight 1 is the table resolving: a `maps_to` this `lower/` lacks names its row."""
    raw = tomllib.loads((HERE / "constructs.toml").read_text(encoding="utf-8"))
    [row] = [r for r in raw["construct"] if r["id"] == "inline.mark"]
    row["maps_to"] = "lower.inline:glow"
    drifted = ConstructTable.parse(raw, NORMALIZERS, source="constructs.toml")
    job, plan, ledger, _ = _job(tmp_path, [section("<p>a <mark>b</mark></p>")])
    compiler = DocxCompiler()
    bundle = ILBundle.load(job / "il", carrier=compiler.carrier, constructs=TABLE)
    with pytest.raises(TargetError, match=r"inline.mark \(lower.inline:glow\)") as caught:
        _pack.build(
            plan,
            bundle,
            ledger,
            carrier=compiler.carrier,
            table=drifted,
            normalizers=NORMALIZERS,
            clock_ns=CLOCK_NS,
            max_units=64,
            provenance=PROVENANCE,
        )
    assert caught.value.code() == "OW_TARGET_CONSTRUCT_UNMAPPED"


@pytest.mark.parametrize(
    "body",
    [
        '<table><tr><td colspan=" 02">a</td></tr></table>',
        '<p lang="EN-gb">x</p>',
    ],
)
def test_a_value_a_normalize_row_would_rewrite_is_refused(tmp_path: Path, body: str) -> None:
    with pytest.raises(TargetError, match="not in canonical form") as caught:
        lowered(tmp_path, section(body))
    assert caught.value.code() == "OW_TARGET_NORMALIZE_AT_COMPILE"


def test_a_slot_the_ledger_lacks_is_unfrozen(tmp_path: Path) -> None:
    missing = slot("ab" * 32)
    with pytest.raises(TargetError) as caught:
        lowered(tmp_path, section(f'<p><img src="{missing}" alt="a"></p>'))
    assert caught.value.code() == "OW_TARGET_ASSET_UNFROZEN"


def test_an_asset_word_cannot_hold_is_refused_by_name(tmp_path: Path) -> None:
    svg = b"<svg xmlns='http://www.w3.org/2000/svg'/>"
    sha = hashlib.sha256(svg).hexdigest()
    with pytest.raises(TargetError, match="PNG, JPEG") as caught:
        lowered(tmp_path, section(f'<p><img src="{slot(sha)}" alt="a"></p>'), assets=[svg])
    assert caught.value.code() == "OW_TARGET_ASSET_FORMAT"


def test_the_plans_roster_is_the_bundles(tmp_path: Path) -> None:
    with pytest.raises(TargetError, match="are not the bundle's") as caught:
        lowered(tmp_path, section("<p>a</p>"), section("<p>b</p>"), plan_units=["u01"])
    assert caught.value.code() == "OW_TARGET_ROSTER_MISMATCH"


def test_the_cards_max_units_is_a_ceiling(tmp_path: Path) -> None:
    job, plan, ledger, _ = _job(tmp_path, [section("<p>a</p>"), section("<p>b</p>")])
    compiler = DocxCompiler()
    bundle = ILBundle.load(job / "il", carrier=compiler.carrier, constructs=TABLE)
    with pytest.raises(TargetError, match="max_units is 1") as caught:
        _pack.build(
            plan,
            bundle,
            ledger,
            carrier=compiler.carrier,
            table=TABLE,
            normalizers={},
            clock_ns=CLOCK_NS,
            max_units=1,
            provenance=PROVENANCE,
        )
    assert caught.value.code() == "OW_TARGET_TOO_MANY_UNITS"


# ---------------------------------------------------------------------------------------------
# the package (09 sections 6 and 12.3)
# ---------------------------------------------------------------------------------------------


def test_the_package_holds_owflows_parts_and_none_of_the_templates_baggage(
    tmp_path: Path,
) -> None:
    out = lowered(tmp_path, section("<p>a</p>"))
    names = set(out.members)
    assert {"word/document.xml", "word/styles.xml", "docProps/core.xml"} <= names
    assert not [n for n in names if "thumbnail" in n or "item1" in n or "Effects" in n]
    assert [n for n in names if n.startswith("customXml/")] == [PROVENANCE]
    assert list(out.members)[:2] == ["[Content_Types].xml", "_rels/.rels"]


def test_no_rsid_and_every_attribute_map_sorted(tmp_path: Path) -> None:
    out = lowered(tmp_path, section('<h1>a</h1><p style="text-align: center">b</p>'))
    for name, payload in out.members.items():
        if not name.endswith((".xml", ".rels")):
            continue
        assert b"rsid" not in payload, name
        for element in etree.fromstring(payload).iter():
            if isinstance(element.tag, str):
                keys = list(element.attrib.keys())
                assert keys == sorted(keys), (name, element.tag)


def test_the_core_properties_are_the_jobs_clock_and_no_host(tmp_path: Path) -> None:
    out = lowered(tmp_path, section("<p>a</p>"))
    core = out.document().core_properties
    when = core.created
    assert when is not None and int(when.timestamp()) == CLOCK_NS // 1_000_000_000
    assert core.modified == when
    assert (core.revision, core.author, core.last_modified_by, core.title) == (1, "", "", "memo")
    app = out.xml("docProps/app.xml")
    assert [(etree.QName(c).localname, c.text) for c in app] == [
        ("Application", "omniweave-target-docx"),
        ("DocSecurity", "0"),
    ]


def test_two_compiles_are_one_package_byte_for_byte(tmp_path: Path) -> None:
    units = (
        section('<h1 data-ow-el="u01/e2">A</h1><ul><li>b</li></ul>', el="u01/e1"),
        section("<table><tr><td>c</td></tr></table>"),
    )
    first = lowered(tmp_path / "one", *units)
    second = lowered(tmp_path / "two", *units)
    assert first.data == second.data


_CHILD = """
import io, sys, zipfile
sys.path.insert(0, sys.argv[1])
from test_docx_lower import lowered, section
from pathlib import Path
from omniweave_core.out.opc import part_digest
out = lowered(Path(sys.argv[2]), section("<h1>A</h1><ol><li>b</li></ol><p>c<br>d</p>"))
print(part_digest(out.members))
"""


def test_the_part_digest_holds_under_another_hash_seed(tmp_path: Path) -> None:
    """09 section 6: `part` is promised; a child with `PYTHONHASHSEED=1` lowers the same parts."""
    here = lowered(tmp_path / "here", section("<h1>A</h1><ol><li>b</li></ol><p>c<br>d</p>"))
    env = {**os.environ, "PYTHONHASHSEED": "1"}
    child = subprocess.run(  # noqa: S603 -- this interpreter, this test's own module
        [sys.executable, "-c", _CHILD, str(Path(__file__).parent), str(tmp_path / "there")],
        capture_output=True,
        text=True,
        env=env,
        check=True,
        timeout=120,
    )
    assert child.stdout.strip() == part_digest(here.members)


# ---------------------------------------------------------------------------------------------
# postflight (09 section 5.3)
# ---------------------------------------------------------------------------------------------


def test_postflight_checks_the_package_and_reports_the_classification(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        section("<p>a</p>"),
        section('<div data-ow-manual="unverifiable_data">x</div>'),
    )
    receipt = out.receipt()
    assert receipt.target == "docx"
    assert receipt.part_digest == part_digest(out.members).removeprefix("sha256:")
    assert receipt.output["sha256"] == hashlib.sha256(out.data).hexdigest()
    assert dict(receipt.units) == {"u01": UnitOutcome.NATIVE, "u02": UnitOutcome.MANUAL_REQUIRED}
    checks = {c.name: (c.value, c.detail) for c in receipt.checks}
    assert checks["unit_count"] == ("passed", "2 sections == unit_count")
    assert checks["page_box"] == ("passed", "a4")
    page = out.document().sections[0]
    assert (page.page_width, page.page_height) == (11906 * 635, 16838 * 635)  # whole twips
    assert {checks[n][0] for n in ("zip_integrity", "internal_relationships", "content_types")} == {
        "passed"
    }
    assert json.loads(receipt.to_bytes())["target"] == "docx"


def test_a_canvas_with_no_page_is_letter_and_a_warning(tmp_path: Path) -> None:
    out = lowered(tmp_path, section("<p>a</p>"), canvas="16:9")
    width = out.document().sections[0].page_width
    assert width is not None and int(width) == 7772400
    [page] = [c for c in out.receipt().checks if c.name == "page_box"]
    assert (page.value, page.detail) == (
        "warning",
        "canvas format '16:9' names no page; the page is letter",
    )


def test_a_package_with_a_unit_missing_is_refused(tmp_path: Path) -> None:
    out = lowered(tmp_path, section("<p>a</p>"), section("<p>b</p>"))
    third = replace(out.plan.units[-1], id="u03", ord=3, il_path="il/u03.html")
    plan = replace(out.plan, unit_count=3, units=(*out.plan.units, third))
    with pytest.raises(TargetError, match="holds 2 units and the plan commissioned 3") as caught:
        DocxCompiler().postflight(str(out.path), plan)
    assert caught.value.code() == "OW_TARGET_UNIT_COUNT_MISMATCH"


def test_lower_job_runs_the_worker_half_over_a_docx_job(tmp_path: Path) -> None:
    """D771's `lower_job()`: the JobDir's digests held, lowered, staged, post-flighted."""
    job, plan, _, store = _job(tmp_path, [section("<p>a</p>")])
    result = lower_job(
        DocxCompiler(),
        Job.load(job),
        job,
        constructs=TABLE,
        suffix=".docx",
        io=make_io(store, tmp_path / "io"),
    )
    artifact, receipt = result.produced
    assert (artifact.kind, receipt.kind) == (ArtifactKind.ARTIFACT, ArtifactKind.RECEIPT)
    body = receipt.inline if receipt.inline is not None else store.open(receipt.blob or "").read()
    assert Receipt.from_json(json.loads(body)).slug == plan.slug


# ---------------------------------------------------------------------------------------------
# the walk: text, blocks, marks
# ---------------------------------------------------------------------------------------------


def test_html_whitespace_collapses_across_runs_and_a_br_breaks_a_line(tmp_path: Path) -> None:
    out = lowered(tmp_path, section("<p>\n  Revenue   <strong>rose</strong>\n 12%.<br> Next  </p>"))
    assert texts(out) == ["Revenue rose 12%.\nNext"]


def test_each_block_takes_words_own_style(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        section(
            "<h1>T</h1><h2>S</h2><h6>D</h6><p>P</p><blockquote><p>Q</p></blockquote>"
            "<pre>x = 1\n  y</pre><hr>"
        ),
    )
    styles = [(p.style.name, p.text) for p in out.document().paragraphs]
    assert styles[:6] == [
        ("Heading 1", "T"),
        ("Heading 2", "S"),
        ("Heading 6", "D"),
        ("Normal", "P"),
        ("Quote", "Q"),
        ("Code Block", "x = 1\n  y"),
    ]
    rule = out.xml().findall(f".//{W}pBdr/{W}bottom")
    assert len(rule) == 1


def test_pre_keeps_its_spaces_but_one_newline_at_each_end(tmp_path: Path) -> None:
    out = lowered(tmp_path, section("<pre>\nx  =  1\n\n  y\n</pre>"))
    assert texts(out) == ["x  =  1\n\n  y"]


@pytest.mark.parametrize(
    ("mark", "element"),
    [
        ("strong", "b"),
        ("em", "i"),
        ("u", "u"),
        ("s", "strike"),
        ("mark", "highlight"),
    ],
)
def test_each_mark_is_its_run_property(tmp_path: Path, mark: str, element: str) -> None:
    out = lowered(tmp_path, section(f"<p>a <{mark}>b</{mark}> c</p>"))
    runs = out.xml().findall(f".//{W}r")
    marked = [r for r in runs if r.find(f"{W}rPr/{W}{element}") is not None]
    assert [_ptext(r) for r in marked] == [" b"]


def test_sub_sup_and_code(tmp_path: Path) -> None:
    out = lowered(tmp_path, section("<p>H<sub>2</sub>O x<sup>2</sup> <code>f()</code></p>"))
    runs = out.xml().findall(f".//{W}r")
    aligned = [r.find(f"{W}rPr/{W}vertAlign").get(f"{W}val") for r in runs
               if r.find(f"{W}rPr/{W}vertAlign") is not None]  # fmt: skip
    assert aligned == ["subscript", "superscript"]
    [code] = [r for r in runs if r.find(f"{W}rPr/{W}rFonts") is not None]
    assert code.find(f"{W}rPr/{W}rFonts").get(f"{W}ascii") == "Courier New"


def test_lang_and_dir_pass_down_and_text_align_is_the_paragraphs(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        section(
            '<blockquote dir="rtl" lang="ar"><p style="text-align: center">a</p><p>b</p>'
            '</blockquote><p>x <span lang="fr">y</span></p>'
        ),
    )
    paragraphs = out.xml().findall(f".//{W}p")
    first = paragraphs[0]
    assert first.find(f"{W}pPr/{W}bidi") is not None
    assert first.find(f"{W}pPr/{W}jc").get(f"{W}val") == "center"
    assert first.find(f".//{W}rPr/{W}rtl") is not None
    assert first.find(f".//{W}rPr/{W}lang").get(f"{W}val") == "ar"
    second = paragraphs[1]
    assert second.find(f"{W}pPr/{W}jc") is None, "text-align is its own paragraph's"
    assert second.find(f"{W}pPr/{W}bidi") is not None, "dir passes down"
    langs = [e.get(f"{W}val") for e in paragraphs[2].iter(f"{W}lang")]
    assert langs == ["fr"]
    assert paragraphs[2].find(f"{W}pPr/{W}bidi") is None


def test_a_manual_region_is_its_own_text_highlighted_with_its_reason_in_a_comment(
    tmp_path: Path,
) -> None:
    """D787: the text is the region's, which R2 holds it to; the reason is the comment's."""
    out = lowered(tmp_path, section('<div data-ow-manual="unverifiable_data">a chart</div>'))
    assert texts(out) == ["a chart"]
    assert len(out.xml().findall(f".//{W}highlight")) == 1
    p = next(p for p in out.xml().iter(f"{W}p") if _ptext(p) == "a chart")
    order = [etree.QName(c).localname for c in p]
    assert order[0] == "commentRangeStart"
    assert order[-2:] == ["commentRangeEnd", "r"]
    assert p[-1].find(f"{W}commentReference") is not None
    [comment] = out.xml("word/comments.xml").findall(f"{W}comment")
    assert _ptext(comment) == "ow:manual unverifiable_data"
    assert (comment.get(f"{W}author"), comment.get(f"{W}initials")) == ("omniweave", "ow")
    assert comment.get(f"{W}date") == "2023-11-14T22:13:20Z", "the job's clock, not the wall's"
    assert "comments+xml" in out.members["[Content_Types].xml"].decode()


# ---------------------------------------------------------------------------------------------
# units, addresses and links
# ---------------------------------------------------------------------------------------------


def test_each_unit_is_a_section_and_the_later_ones_continue_the_page(tmp_path: Path) -> None:
    out = lowered(tmp_path, section("<p>a</p>"), section("<p>b</p>"), section("<p>c</p>"))
    sections = out.xml().findall(f".//{W}sectPr")
    assert len(sections) == 3
    kinds = [s.find(f"{W}type") for s in sections]
    assert kinds[0] is None
    assert [k.get(f"{W}val") for k in kinds[1:] if k is not None] == ["continuous", "continuous"]


def test_every_address_is_a_bookmark_inside_paragraphs(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        section(
            '<h1 data-ow-el="u01/e2">A</h1><p data-ow-el="u01/e3">b <strong data-ow-el="u01/e4">'
            'c</strong></p><ul data-ow-el="u01/e5"><li>d</li></ul>',
            el="u01/e1",
        ),
    )
    root = out.xml()
    starts = {s.get(f"{W}name"): s.get(f"{W}id") for s in root.iter(f"{W}bookmarkStart")}
    assert set(starts) == {f"ow_u01_e{n}" for n in range(1, 6)}
    ends = [e.get(f"{W}id") for e in root.iter(f"{W}bookmarkEnd")]
    assert sorted(ends) == sorted(starts.values())
    assert len(set(starts.values())) == 5
    for mark in [*root.iter(f"{W}bookmarkStart"), *root.iter(f"{W}bookmarkEnd")]:
        assert etree.QName(mark.getparent()).localname == "p"


def test_links_are_anchors_or_external_relationships(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        section(
            '<h1 data-ow-el="u01/e2">A</h1><p>see <a href="#u01/e2">above</a> or '
            '<a href="mailto:x@y.invalid">mail</a>.</p>'
        ),
    )
    links = out.xml().findall(f".//{W}hyperlink")
    assert links[0].get(f"{W}anchor") == "ow_u01_e2"
    assert [_ptext(link) for link in links] == ["above", "mail"]
    rels = out.members["word/_rels/document.xml.rels"].decode()
    assert 'Target="mailto:x@y.invalid"' in rels
    assert 'TargetMode="External"' in rels
    assert texts(out)[-1] == "see above or mail."


# ---------------------------------------------------------------------------------------------
# lists and tables
# ---------------------------------------------------------------------------------------------


def _numpr(p: Any) -> tuple[str, str] | None:
    numpr = p.find(f"{W}pPr/{W}numPr")
    if numpr is None:
        return None
    return numpr.find(f"{W}ilvl").get(f"{W}val"), numpr.find(f"{W}numId").get(f"{W}val")


def test_lists_nest_by_level_and_restart_where_they_begin(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        section(
            '<ul><li>a<ol start="3"><li>b</li></ol></li><li><p>c</p><p>d</p></li></ul>'
            "<ol><li>e</li></ol>"
        ),
    )
    paragraphs = [p for p in out.xml().iter(f"{W}p") if _ptext(p)]
    levels = [(_ptext(p), _numpr(p)) for p in paragraphs]
    assert [(t, n[0] if n else None) for t, n in levels] == [
        ("a", "0"),
        ("b", "1"),
        ("c", "0"),
        ("d", None),
        ("e", "0"),
    ]
    ids = {t: n[1] for t, n in levels if n}
    assert ids["a"] == ids["c"] != ids["b"] != ids["e"]
    numbering = out.xml("word/numbering.xml")
    [num] = [n for n in numbering.iter(f"{W}num") if n.get(f"{W}numId") == ids["b"]]
    override = num.find(f"{W}lvlOverride")
    assert override.get(f"{W}ilvl") == "1"
    assert override.find(f"{W}startOverride").get(f"{W}val") == "3"
    [d] = [p for p in paragraphs if _ptext(p) == "d"]
    assert d.find(f"{W}pPr/{W}ind").get(f"{W}left") == "720"


def test_a_definition_list_is_a_bold_term_and_an_indented_description(tmp_path: Path) -> None:
    out = lowered(tmp_path, section("<dl><dt>Risk</dt><dd>low</dd></dl>"))
    term, desc = [p for p in out.xml().iter(f"{W}p") if _ptext(p)]
    assert term.find(f"{W}pPr/{W}keepNext") is not None
    assert term.find(f".//{W}b") is not None
    assert desc.find(f"{W}pPr/{W}ind").get(f"{W}left") == "720"


def test_a_table_is_laid_out_merged_and_headed(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        section(
            '<table><caption>Q3</caption><thead><tr><th colspan="2">Region</th></tr></thead>'
            '<tbody><tr><td rowspan="2">EU</td><td>1</td></tr><tr><td>2</td></tr></tbody></table>'
        ),
    )
    body = out.xml().find(f"{W}body")
    caption, table_ = body[0], body[1]
    assert etree.QName(table_).localname == "tbl"
    assert caption.find(f"{W}pPr/{W}pStyle").get(f"{W}val") == "Caption"
    rows = table_.findall(f"{W}tr")
    assert rows[0].find(f"{W}trPr/{W}tblHeader") is not None
    assert rows[0].find(f".//{W}gridSpan").get(f"{W}val") == "2"
    assert rows[0].find(f".//{W}b") is not None
    merges = [tc.find(f"{W}tcPr/{W}vMerge") for tc in table_.iter(f"{W}tc")]
    assert [m.get(f"{W}val") for m in merges if m is not None] == ["restart", None]
    assert [_ptext(tc) for tc in table_.iter(f"{W}tc")] == ["Region", "EU", "1", "", "2"]


def test_a_cell_keeps_one_paragraph_once_it_has_content(tmp_path: Path) -> None:
    out = lowered(tmp_path, section("<table><tr><td><p>a</p><p>b</p></td></tr></table>"))
    [tc] = out.xml().iter(f"{W}tc")
    assert [_ptext(p) for p in tc.findall(f"{W}p")] == ["a", "b"]


# ---------------------------------------------------------------------------------------------
# pictures
# ---------------------------------------------------------------------------------------------


def test_a_picture_fits_the_text_column_in_whole_emus(tmp_path: Path) -> None:
    wide, small = (slot(hashlib.sha256(b).hexdigest()) for b in (WIDE, SMALL))
    out = lowered(
        tmp_path,
        section(
            f'<figure><img src="{wide}" alt="a wide chart"><figcaption>C</figcaption></figure>'
            f'<p><img src="{small}" alt="small"> and <img src="{small}" alt="again"></p>'
        ),
        assets=[WIDE, SMALL],
    )
    page = out.document().sections[0]
    column = int(page.page_width or 0) - int(page.left_margin or 0) - int(page.right_margin or 0)
    extents = [(int(e.get("cx")), int(e.get("cy"))) for e in out.xml().iter(f"{WP}extent")]
    native = Image.from_blob(WIDE)
    assert extents[0] == (column, int(native.height) * column // int(native.width))
    assert extents[1] == extents[2] == (508000, 254000)
    descr = [e.get("descr") for e in out.xml().iter(f"{WP}docPr")]
    assert descr == ["a wide chart", "small", "again"]
    media = [n for n in out.members if n.startswith("word/media/")]
    assert len(media) == 2


# ---------------------------------------------------------------------------------------------
# D783: notes
# ---------------------------------------------------------------------------------------------


def _notes(out: Lowered, kind: str = "footnote") -> list[Any]:
    root = out.xml(f"word/{kind}s.xml")
    return [n for n in root.findall(f"{W}{kind}") if n.get(f"{W}type") is None]


def test_a_footnote_is_a_reference_in_the_text_and_a_body_in_its_part(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        section(
            '<p>Rose.<a data-ow-note="n1"></a> Held.</p>'
            '<aside data-ow-footnote="n1"><p>Audited.</p></aside>'
        ),
    )
    assert texts(out) == ["Rose. Held."]
    [marker] = out.xml().iter(f"{W}footnoteReference")
    assert marker.get(f"{W}id") == "1"
    [note] = _notes(out)
    assert (note.get(f"{W}id"), _ptext(note)) == ("1", " Audited.")
    assert note.find(f".//{W}footnoteRef") is not None
    separators = [n.get(f"{W}type") for n in out.xml("word/footnotes.xml").findall(f"{W}footnote")]
    assert separators[:2] == ["separator", "continuationSeparator"]
    types = out.members["[Content_Types].xml"].decode()
    assert "wordprocessingml.footnotes+xml" in types
    assert "word/endnotes.xml" not in out.members


def test_notes_are_numbered_in_reference_order_and_each_part_counts_its_own(
    tmp_path: Path,
) -> None:
    out = lowered(
        tmp_path,
        section(
            '<aside data-ow-footnote="b"><p>second</p></aside>'
            '<aside data-ow-endnote="c"><p>end</p></aside>'
            '<aside data-ow-footnote="a"><p>first</p></aside>'
            '<p>x<a data-ow-note="a"></a> y<a data-ow-note="b"></a> z<a data-ow-note="c"></a></p>'
        ),
    )
    assert [(n.get(f"{W}id"), _ptext(n).strip()) for n in _notes(out)] == [
        ("1", "first"),
        ("2", "second"),
    ]
    assert [(n.get(f"{W}id"), _ptext(n).strip()) for n in _notes(out, "endnote")] == [("1", "end")]
    refs = [
        (etree.QName(r).localname, r.get(f"{W}id"))
        for r in out.xml().iter(f"{W}footnoteReference", f"{W}endnoteReference")
    ]
    assert refs == [
        ("footnoteReference", "1"),
        ("footnoteReference", "2"),
        ("endnoteReference", "1"),
    ]


def test_a_note_is_its_own_story(tmp_path: Path) -> None:
    """Nothing the reference stands inside reaches the body: no list, no bold, no alignment."""
    out = lowered(
        tmp_path,
        section(
            '<ul><li style="text-align: center"><strong>x<a data-ow-note="n1"></a></strong>'
            ' y</li></ul><aside data-ow-footnote="n1"><p>body</p><table><tr><td>t</td></tr>'
            "</table></aside>"
        ),
    )
    [note] = _notes(out)
    first = note.find(f"{W}p")
    assert first.find(f"{W}pPr/{W}pStyle").get(f"{W}val") == "FootnoteText"
    assert first.find(f"{W}pPr/{W}numPr") is None
    assert first.find(f"{W}pPr/{W}jc") is None
    assert first.find(f".//{W}b") is None
    assert note.find(f"{W}tbl") is not None
    assert etree.QName(note[-1]).localname == "p"
    [item] = [p for p in out.xml().iter(f"{W}p") if _ptext(p).strip().startswith("x")]
    assert item.find(f"{W}pPr/{W}numPr") is not None
    assert texts(out)[0] == "x y"


def test_a_note_bodys_address_is_a_bookmark_inside_it(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        section(
            '<p>x<a data-ow-note="n1"></a></p>'
            '<aside data-ow-footnote="n1" data-ow-el="u01/e9"><p>body</p></aside>'
        ),
    )
    [note] = _notes(out)
    [start] = note.iter(f"{W}bookmarkStart")
    assert start.get(f"{W}name") == "ow_u01_e9"


def test_the_note_styles_are_words_own_ids(tmp_path: Path) -> None:
    out = lowered(tmp_path, section("<p>a</p>"))
    ids = {s.get(f"{W}styleId") for s in out.xml("word/styles.xml").iter(f"{W}style")}
    assert {"FootnoteText", "FootnoteReference", "EndnoteText", "EndnoteReference"} <= ids


# ---------------------------------------------------------------------------------------------
# D783: the provenance part
# ---------------------------------------------------------------------------------------------


def _provenance(out: Lowered) -> dict[str, Any]:
    return json.loads(out.members[PROVENANCE])


def test_the_provenance_part_is_written_even_with_nothing_cited(tmp_path: Path) -> None:
    out = lowered(tmp_path, section("<p>a</p>"))
    assert _provenance(out) == {
        "artifact_gen": 1,
        "cites": [],
        "corpus": "acme",
        "els": {},
        "schema": "omniweave.out-provenance/1",
        "slug": "memo",
    }
    rels = out.members["_rels/.rels"].decode()
    assert 'Type="urn:omniweave:relationship:provenance"' in rels
    assert 'Target="customXml/omniweave-provenance.json"' in rels
    assert "application/json" in out.members["[Content_Types].xml"].decode()
    [row] = [c for c in out.receipt().checks if c.name == "citations_channel"]
    assert (row.value, row.detail) == ("passed", "embedded: 0 cite row(s)")


def test_each_cite_is_a_row_with_its_pinned_generation(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        section(
            '<p data-ow-el="u01/e2" data-ow-cite="d7#1 acme:d7#2">x</p>'
            '<p data-ow-el="u01/e3" data-ow-cite="d7#1" data-ow-cite-kind="quote"'
            ' data-ow-cite-span="0:7">y</p>'
            '<p data-ow-el="u01/e4" data-ow-cite="d7#1 d7#2" data-ow-cite-span="0:1">z</p>'
        ),
        cites={"d7#1": 3, "d7#2": 5},
    )
    assert _provenance(out)["cites"] == [
        {"cite": "d7#1", "doc_gen": 3, "el": "u01/e2", "kind": None, "span": None},
        {"cite": "d7#2", "doc_gen": 5, "el": "u01/e2", "kind": None, "span": None},
        {"cite": "d7#1", "doc_gen": 3, "el": "u01/e3", "kind": "quote", "span": [0, 7]},
        {"cite": "d7#1", "doc_gen": 3, "el": "u01/e4", "kind": None, "span": None},
        {"cite": "d7#2", "doc_gen": 5, "el": "u01/e4", "kind": None, "span": None},
    ]
    [row] = [c for c in out.receipt().checks if c.name == "citations_channel"]
    assert row.detail == "embedded: 5 cite row(s)"
    assert texts(out) == ["x", "y", "z"]


@pytest.mark.parametrize(
    ("body", "code", "said"),
    [
        ('<p data-ow-cite="d7#1">x</p>', "OW_TARGET_IL_MALFORMED", "carries no data-ow-el"),
        ('<p data-ow-el="u01/e2" data-ow-cite="d7#9">x</p>', "OW_TARGET_CITE_UNRESOLVED", "d7#9"),
        (
            '<p data-ow-el="u01/e2" data-ow-cite="d7#1" data-ow-cite-kind="rumour">x</p>',
            "OW_TARGET_IL_MALFORMED",
            "is not one of",
        ),
    ],
)
def test_a_cite_the_part_cannot_state_is_refused(
    tmp_path: Path, body: str, code: str, said: str
) -> None:
    with pytest.raises(TargetError, match=said) as caught:
        lowered(tmp_path, section(body), cites={"d7#1": 3})
    assert caught.value.code() == code


def test_a_package_without_its_provenance_part_is_a_failed_row(tmp_path: Path) -> None:
    out = lowered(tmp_path, section("<p>a</p>"))
    members = {k: v for k, v in out.members.items() if k != PROVENANCE}
    rels = members["_rels/.rels"].decode()
    start = rels.index("<Relationship", rels.index("urn:omniweave") - 200)
    members["_rels/.rels"] = (rels[:start] + rels[rels.index("/>", start) + 2 :]).encode()
    stripped = tmp_path / "stripped.docx"
    stripped.write_bytes(write_opc(members))
    [row] = [c for c in DocxCompiler().postflight(str(stripped), out.plan).checks
             if c.name == "citations_channel"]  # fmt: skip
    assert (row.value, row.detail) == ("failed", f"no {PROVENANCE} in the package")


def test_the_part_maps_every_address_to_its_construct(tmp_path: Path) -> None:
    """D786: a bookmark cannot say whether `u01/e1` is the section or the heading it opens with;
    the construct can, through `kind_map.toml`, for `parse.docx.native`."""
    out = lowered(
        tmp_path,
        section(
            '<h1 data-ow-el="u01/e2">A</h1><p data-ow-el="u01/e3">b <strong data-ow-el="u01/e4">'
            'c</strong><a data-ow-note="n1"></a></p><table data-ow-el="u01/e5"><tr>'
            '<td data-ow-el="u01/e6">d</td></tr></table>'
            '<aside data-ow-footnote="n1" data-ow-el="u01/e7"><p>e</p></aside>',
            el="u01/e1",
        ),
    )
    assert _provenance(out)["els"] == {
        "u01/e1": "ow.unit",
        "u01/e2": "block.h1",
        "u01/e3": "block.p",
        "u01/e4": "inline.strong",
        "u01/e5": "table.table",
        "u01/e6": "table.td",
        "u01/e7": "note.footnote",
    }


def _unzipped(raw: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(raw)) as zf:
        return {info.filename: zf.read(info) for info in zf.infolist()}


def test_the_collected_members_are_what_document_save_zips(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """D887: `_members()` walks python-docx's private writer steps; this is what notices a
    python-docx that moves them. Then the package builds under 09 section 6's
    `SOURCE_DATE_EPOCH=0`, which 3.14's `zipfile` turned into a `struct.error` inside `save()`."""
    document = docx.Document()
    document.add_paragraph("one")
    monkeypatch.delenv("SOURCE_DATE_EPOCH", raising=False)
    saved = io.BytesIO()
    document.save(saved)
    assert _pack._members(document) == _unzipped(saved.getvalue())
    monkeypatch.setenv("SOURCE_DATE_EPOCH", "0")
    assert set(_unzipped(_pack._package(document))) == set(_unzipped(saved.getvalue()))
