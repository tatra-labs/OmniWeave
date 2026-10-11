"""The notes parts: `word/footnotes.xml` and `word/endnotes.xml`, and a container over one note.

A note body is written through the same walk as the document, into a `w:footnote` or
`w:endnote` instead of the body, so a `Story` gives that element the two methods the walk's
scopes call, `add_paragraph()` and `add_table()`, as python-docx's `_Cell` gives a `w:tc`.
"""

from __future__ import annotations

from typing import TYPE_CHECKING, Any, Final

from docx.opc.constants import CONTENT_TYPE, RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.shared import Emu
from docx.table import Table
from docx.text.paragraph import Paragraph

if TYPE_CHECKING:
    from docx.document import Document
    from docx.opc.part import Part

__all__ = ["KINDS", "REFERENCE_STYLE", "STYLE_IDS", "TEXT_STYLE", "Notes", "Story", "mark"]

KINDS: Final = ("footnote", "endnote")
TEXT_STYLE: Final = {"footnote": "footnote text", "endnote": "endnote text"}
REFERENCE_STYLE: Final = {"footnote": "footnote reference", "endnote": "endnote reference"}
STYLE_IDS: Final = {
    "footnote text": "FootnoteText",
    "footnote reference": "FootnoteReference",
    "endnote text": "EndnoteText",
    "endnote reference": "EndnoteReference",
}
"""Word's own style ids for its four note styles, which python-docx's template does not define."""
_PART: Final = {
    "footnote": (CONTENT_TYPE.WML_FOOTNOTES, RELATIONSHIP_TYPE.FOOTNOTES),
    "endnote": (CONTENT_TYPE.WML_ENDNOTES, RELATIONSHIP_TYPE.ENDNOTES),
}
_SEPARATORS: Final = (("-1", "separator"), ("0", "continuationSeparator"))
"""The two notes Word draws the rule above the notes area with, ids -1 and 0."""


class Story:
    """A block container over one note element, the walk's scope inside a note."""

    def __init__(self, element: Any, part: Part, width: int) -> None:
        self._element = element
        self._part = part
        self._width = width

    @property
    def part(self) -> Part:
        return self._part

    def add_paragraph(self, style: str | None = None) -> Paragraph:
        p = OxmlElement("w:p")
        self._element.append(p)
        paragraph = Paragraph(p, self)  # type: ignore[arg-type]
        if style:
            paragraph.style = style
        return paragraph

    def add_table(self, rows: int, cols: int) -> Table:
        """A table, then the empty paragraph a story must end with, as python-docx's cell adds."""
        tbl = CT_Tbl.new_tbl(rows, cols, Emu(self._width))
        self._element.append(tbl)
        self._element.append(OxmlElement("w:p"))
        return Table(tbl, self)  # type: ignore[arg-type]


class Notes:
    """Both notes parts, made on first use, and where each note id was placed."""

    def __init__(self) -> None:
        self._roots: dict[str, Any] = {}
        self._counts: dict[str, int] = dict.fromkeys(KINDS, 0)
        self.placed: dict[str, tuple[str, int]] = {}
        """IL note id to `(kind, w:id)`: what a reference names."""

    def new(self, kind: str) -> tuple[int, Any]:
        """The next `w:footnote` or `w:endnote`, numbered in reference order from 1."""
        root = self._roots.get(kind)
        if root is None:
            root = self._roots[kind] = _root(kind)
        self._counts[kind] += 1
        note = OxmlElement(f"w:{kind}")
        note.set(qn("w:id"), str(self._counts[kind]))
        root.append(note)
        return self._counts[kind], note

    def attach(self, document: Document) -> None:
        """Each part that holds a note, related from the document in a fixed order."""
        for kind in KINDS:
            root = self._roots.get(kind)
            if root is None:
                continue
            content_type, reltype = _PART[kind]
            package = document.part.package
            if package is None:  # pragma: no cover -- a loaded document always has one
                continue
            part = XmlPart(PackURI(f"/word/{kind}s.xml"), content_type, root, package)
            document.part.relate_to(part, reltype)


def mark(note: Any, kind: str) -> None:
    """The note's own number at the start of its first paragraph, then a space."""
    first = note.find(qn("w:p"))
    if first is None:
        first = OxmlElement("w:p")
        style = OxmlElement("w:pStyle")
        style.set(qn("w:val"), STYLE_IDS[TEXT_STYLE[kind]])
        ppr = OxmlElement("w:pPr")
        ppr.append(style)
        first.append(ppr)
        note.append(first)
    number = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rstyle = OxmlElement("w:rStyle")
    rstyle.set(qn("w:val"), STYLE_IDS[REFERENCE_STYLE[kind]])
    rpr.append(rstyle)
    number.append(rpr)
    number.append(OxmlElement(f"w:{kind}Ref"))
    space = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.set(qn("xml:space"), "preserve")
    text.text = " "
    space.append(text)
    ppr = first.find(qn("w:pPr"))
    if ppr is not None:
        ppr.addnext(number)
    else:
        first.insert(0, number)
    number.addnext(space)


def _root(kind: str) -> Any:
    root = OxmlElement(f"w:{kind}s")
    for number, role in _SEPARATORS:
        note = OxmlElement(f"w:{kind}")
        note.set(qn("w:type"), role)
        note.set(qn("w:id"), number)
        p = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        spacing = OxmlElement("w:spacing")
        for name, value in (("after", "0"), ("line", "240"), ("lineRule", "auto")):
            spacing.set(qn(f"w:{name}"), value)
        ppr.append(spacing)
        p.append(ppr)
        r = OxmlElement("w:r")
        r.append(OxmlElement(f"w:{role}"))
        p.append(r)
        note.append(p)
        root.append(note)
    return root
