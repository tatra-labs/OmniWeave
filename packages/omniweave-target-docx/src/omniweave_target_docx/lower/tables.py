"""Tables: HTML's table model laid out on a grid, merged, then filled through the walk.

The gate held every row group to a grid its cells cover exactly (`docx.table.topology@1`), so
the layout here assumes it and does not check it again. Header rows repeat on each page
(`w:tblHeader`), a header cell's runs are bold, and the caption is a `Caption` paragraph before
the table, where Word's own captions go.
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from typing import TYPE_CHECKING, Any, Final

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

if TYPE_CHECKING:
    from docx.table import Table, _Cell
    from omniweave_core.out.carrier import CarrierNode

    from omniweave_target_docx.lower._emit import Emitter

__all__ = [
    "body_rows",
    "caption",
    "cell",
    "grid_span",
    "header_cell",
    "header_rows",
    "row",
    "table",
    "vertical_merge",
]

STYLE: Final = "Table Grid"
_GROUPS: Final = frozenset({"html|thead", "html|tbody"})
_CELLS: Final = frozenset({"html|th", "html|td"})
_P: Final = qn("w:p")
_PPR: Final = qn("w:pPr")


@dataclass(frozen=True, slots=True)
class _Placed:
    node: CarrierNode
    row: int
    col: int
    rows: int
    cols: int


@dataclass(frozen=True, slots=True)
class _Group:
    node: CarrierNode | None
    """`thead` or `tbody`, or `None` for rows written straight under `table`."""
    header: bool
    rows: tuple[CarrierNode, ...]
    first: int
    """The grid row the group's first row lands on."""


def table(em: Emitter, node: CarrierNode) -> None:
    """The caption, then the grid: laid out, merged, then each group filled in order."""
    groups = _groups(node)
    placed, height, width = _layout(groups)
    for child in node.children():
        if child.tag == "html|caption":
            em.block(child)
    em.end_paragraph()
    grid: Table = em.scopes[-1].container.add_table(rows=max(height, 1), cols=max(width, 1))
    grid.style = STYLE
    for p in placed:
        if p.rows > 1:
            vertical_merge(grid, p.row, p.col, p.rows, p.cols)
        elif p.cols > 1:
            grid_span(grid, p.row, p.col, p.cols)
    by_row: dict[int, list[_Placed]] = {}
    for p in placed:
        by_row.setdefault(p.row, []).append(p)
    for group in groups:
        (header_rows if group.header else body_rows)(em, grid, group, by_row)


def caption(em: Emitter, node: CarrierNode) -> None:
    em.new_paragraph("Caption")
    em.inline_children(node)


def header_rows(em: Emitter, grid: Table, group: Any, by_row: dict[int, Any]) -> None:
    """`thead`: rows Word repeats at the top of every page the table crosses."""
    for offset, tr in enumerate(group.rows):
        at = group.first + offset
        flag = OxmlElement("w:tblHeader")
        grid.rows[at]._tr.get_or_add_trPr().append(flag)
        row(em, grid, tr, by_row.get(at, []))
    last = group.first + len(group.rows) - 1
    _address(em, grid, group.node, by_row, top=group.first, bottom=last)


def body_rows(em: Emitter, grid: Table, group: Any, by_row: dict[int, Any]) -> None:
    """`tbody`, or the rows written straight under `table`."""
    for offset, tr in enumerate(group.rows):
        row(em, grid, tr, by_row.get(group.first + offset, []))
    last = group.first + len(group.rows) - 1
    _address(em, grid, group.node, by_row, top=group.first, bottom=last)


def row(em: Emitter, grid: Table, tr: CarrierNode, cells: Sequence[Any]) -> None:
    """One `tr`: each cell it places, filled, and its `data-ow-el` across them."""
    for p in cells:
        target = grid.cell(p.row, p.col)
        (header_cell if p.node.tag == "html|th" else cell)(em, p.node, target)
    if cells and tr.el:
        _span_bookmark(em, grid.cell(cells[0].row, cells[0].col), _last(grid, cells), tr.el)


def header_cell(em: Emitter, node: CarrierNode, target: _Cell) -> None:
    with em.marked(bold=True):
        cell(em, node, target)


def cell(em: Emitter, node: CarrierNode, target: _Cell) -> None:
    """The cell's contents through the walk; its own empty paragraph goes once it has others."""
    tc = target._tc
    own = tc.find(_P)
    with em.scope(target, element=tc):
        em.addressed(node, lambda: em.flow(node))
    if own is not None and len(own) == 0 and len(tc.findall(_P)) + len(tc.findall(qn("w:tbl"))) > 1:
        tc.remove(own)


def grid_span(grid: Table, at_row: int, at_col: int, cols: int) -> None:
    """`colspan`: one cell across `cols` grid columns (`w:gridSpan`)."""
    grid.cell(at_row, at_col).merge(grid.cell(at_row, at_col + cols - 1))


def vertical_merge(grid: Table, at_row: int, at_col: int, rows: int, cols: int) -> None:
    """`rowspan`: one cell down `rows` grid rows (`w:vMerge`), and across `cols` with it."""
    grid.cell(at_row, at_col).merge(grid.cell(at_row + rows - 1, at_col + cols - 1))


def _groups(node: CarrierNode) -> list[_Group]:
    """Row groups in document order; consecutive bare rows are one group."""
    out: list[_Group] = []
    bare: list[CarrierNode] = []
    first = 0

    def close() -> None:
        nonlocal first, bare
        if bare:
            out.append(_Group(None, header=False, rows=tuple(bare), first=first))
            first += len(bare)
            bare = []

    for child in node.children():
        if child.tag == "html|tr":
            bare.append(child)
        elif child.tag in _GROUPS:
            close()
            rows = tuple(r for r in child.children() if r.tag == "html|tr")
            out.append(_Group(child, header=child.tag == "html|thead", rows=rows, first=first))
            first += len(rows)
    close()
    return out


def _layout(groups: Sequence[_Group]) -> tuple[list[_Placed], int, int]:
    """HTML's table model, one group at a time: each cell's grid origin and span."""
    placed: list[_Placed] = []
    width = 0
    for group in groups:
        taken: set[tuple[int, int]] = set()
        for offset, tr in enumerate(group.rows):
            c = 0
            for td in tr.children():
                if td.tag not in _CELLS:
                    continue
                while (offset, c) in taken:
                    c += 1
                cols, rows = _span(td, "colspan"), _span(td, "rowspan")
                rows = min(rows, len(group.rows) - offset)
                taken |= {(offset + dr, c + dc) for dr in range(rows) for dc in range(cols)}
                placed.append(_Placed(td, group.first + offset, c, rows, cols))
                c += cols
        width = max([width, *(col + 1 for _, col in taken)])
    height = sum(len(g.rows) for g in groups)
    return placed, height, width


def _span(node: CarrierNode, name: str) -> int:
    value = (node.attr(name) or "1").strip().removeprefix("+")
    return int(value) if value.isdecimal() and int(value) > 0 else 1


def _last(grid: Table, cells: Sequence[_Placed]) -> _Cell:
    last = cells[-1]
    return grid.cell(last.row, last.col)


def _address(
    em: Emitter,
    grid: Table,
    node: CarrierNode | None,
    by_row: dict[int, list[_Placed]],
    *,
    top: int,
    bottom: int,
) -> None:
    """A row group's `data-ow-el`, from its first cell's text to its last's."""
    if node is None or not node.el or not by_row.get(top) or not by_row.get(bottom):
        return
    first, last = by_row[top][0], by_row[bottom][-1]
    _span_bookmark(em, grid.cell(first.row, first.col), grid.cell(last.row, last.col), node.el)


def _span_bookmark(em: Emitter, first: _Cell, last: _Cell, el: str) -> None:
    """A bookmark starting in `first`'s first paragraph and ending in `last`'s last."""
    start, end = em.bookmark_pair(el)
    opening = first._tc.find(_P)
    closing = last._tc.findall(_P)[-1]
    ppr = opening.find(_PPR)
    if ppr is not None:
        ppr.addnext(start)
    else:
        opening.insert(0, start)
    closing.append(end)
