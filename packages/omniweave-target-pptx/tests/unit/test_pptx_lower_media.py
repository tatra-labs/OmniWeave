"""`compile.pptx.native`'s pictures: an `<image>` fitted, cropped and clipped. D800.

Each picture is a PNG of `png(width, height)` frozen into the ledger (tests/lowerkit.py), lowered
and read back from the package and with python-pptx. The shown image, what is seen and the crop
follow `lower.media`'s three boxes, so each expected value is worked by hand.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from lowerkit import (  # type: ignore[import-not-found]  # tests/lowerkit.py
    A,
    P,
    box,
    lowered,
    outline,
    picture_slot,
    png,
    slide,
    turn,
    u,
)
from omniweave_core.errors import TargetError
from omniweave_target_pptx.driver import PptxCompiler
from omniweave_target_pptx.lower.defs import reference
from omniweave_target_pptx.lower.media import aspect
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

WIDE = png(200, 100)
SQUARE = png(100, 100)


def pictured(tmp_path: Path, attrs: str, *, blob: bytes = WIDE, defs: str = "") -> Any:
    """The first picture of a slide holding one `<image>` of `blob` with `attrs`."""
    image = f'<image href="{picture_slot(blob)}" {attrs}/>'
    out = lowered(tmp_path, slide(f"<defs>{defs}</defs>{image}"), assets=[blob])
    return out.shapes()[0]


def crop(shape: Any) -> dict[str, str]:
    found = shape.find(f"{P}blipFill/{A}srcRect")
    return {} if found is None else dict(found.attrib)


def geometry(shape: Any) -> list[str]:
    return outline(shape.find(f"{P}spPr")[1])


# ---------------------------------------------------------------------------------------------
# the picture
# ---------------------------------------------------------------------------------------------


def test_a_picture_is_a_pic_in_drawingmls_order(tmp_path: Path) -> None:
    shape = pictured(tmp_path, 'data-ow-el="u01/e2" x="10" y="20" width="100" height="100"')
    assert [c for c in outline(shape) if not c.startswith("blip ")] == [
        "pic",
        "nvPicPr",
        "cNvPr id=2 name=u01/e2",
        "cNvPicPr",
        "picLocks noChangeAspect=1",
        "nvPr",
        "blipFill",
        "stretch",
        "fillRect",
        "spPr",
        "xfrm",
        f"off x={u(10)} y={u(45)}",
        f"ext cx={u(100)} cy={u(50)}",
        "prstGeom prst=rect",
        "avLst",
    ]
    assert shape.find(f"{P}blipFill/{A}blip").get(qn("r:embed")) == "rId2"


def test_python_pptx_reads_the_picture_and_its_part_back(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide(f'<image href="{picture_slot(WIDE)}" width="100" height="100"/>'),
        assets=[WIDE],
    )
    [picture] = out.deck().slides[0].shapes
    assert (picture.shape_type, picture.image.size, picture.image.blob) == (
        MSO_SHAPE_TYPE.PICTURE,
        (200, 100),
        WIDE,
    )
    assert [n for n in out.members if n.startswith("ppt/media/")] == ["ppt/media/image1.png"]


def test_one_picture_twice_is_one_part_and_one_relationship(tmp_path: Path) -> None:
    image = f'<image href="{picture_slot(WIDE)}" width="10" height="10"/>'
    out = lowered(tmp_path, slide(image + image), assets=[WIDE])
    assert [n for n in out.members if n.startswith("ppt/media/")] == ["ppt/media/image1.png"]
    assert {s.find(f"{P}blipFill/{A}blip").get(qn("r:embed")) for s in out.shapes()} == {"rId2"}


# ---------------------------------------------------------------------------------------------
# preserveAspectRatio
# ---------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("ratio", "expected", "cropped"),
    [
        ("", (u(10), u(45), u(100), u(50)), {}),
        ('preserveAspectRatio="xMidYMid meet"', (u(10), u(45), u(100), u(50)), {}),
        ('preserveAspectRatio="xMinYMax"', (u(10), u(70), u(100), u(50)), {}),
        ('preserveAspectRatio="xMidYMin meet"', (u(10), u(20), u(100), u(50)), {}),
        ('preserveAspectRatio="none"', (u(10), u(20), u(100), u(100)), {}),
        (
            'preserveAspectRatio="xMidYMid slice"',
            (u(10), u(20), u(100), u(100)),
            {"l": "25000", "r": "25000"},
        ),
        ('preserveAspectRatio="xMaxYMin slice"', (u(10), u(20), u(100), u(100)), {"l": "50000"}),
        ('preserveAspectRatio="xMinYMid slice"', (u(10), u(20), u(100), u(100)), {"r": "50000"}),
    ],
)
def test_the_aspect_fits_the_picture_and_slice_crops_what_spills(
    tmp_path: Path, ratio: str, expected: tuple[int, int, int, int], cropped: dict[str, str]
) -> None:
    shape = pictured(tmp_path, f'x="10" y="20" width="100" height="100" {ratio}')
    assert box(shape) == expected
    assert crop(shape) == cropped


def test_an_aspect_word_the_gate_refuses_is_the_initial_meet() -> None:
    root = (
        PptxCompiler()
        .carrier(
            "il/u01.svg",
            slide(
                '<image preserveAspectRatio="xMaxYMax slice"/><image preserveAspectRatio="bogus"/>'
                '<image href="x"/>'
            ),
        )
        .root()
    )
    first, second, third = root.children()
    assert (aspect(first), aspect(second)) == ((1.0, 1.0, "slice"), (0.5, 0.5, "meet"))
    assert (reference(third), reference(first)) == ("x", "")


# ---------------------------------------------------------------------------------------------
# what else a picture takes
# ---------------------------------------------------------------------------------------------


def test_a_pictures_opacity_is_its_blips_alpha(tmp_path: Path) -> None:
    shape = pictured(tmp_path, 'width="10" height="10" opacity="0.5"')
    blip = shape.find(f"{P}blipFill/{A}blip")
    assert [c for c in outline(blip) if not c.startswith("blip ")] == ["alphaModFix amt=50000"]


def test_a_turned_picture_turns_its_frame(tmp_path: Path) -> None:
    shape = pictured(tmp_path, 'width="100" height="100" transform="rotate(90 50 50)"')
    assert turn(shape) == ("5400000", None, None)


def test_a_pictures_filter_is_its_effect_list(tmp_path: Path) -> None:
    shape = pictured(
        tmp_path,
        'width="10" height="10" filter="url(#f)"',
        defs='<filter id="f"><feDropShadow dx="3" dy="0" stdDeviation="1"/></filter>',
    )
    body = shape.find(f"{P}spPr")
    assert [c.tag.rpartition("}")[2] for c in body] == ["xfrm", "prstGeom", "effectLst"]
    assert outline(body[2])[1] == f"outerShdw blurRad={u(2)} dir=0 dist={u(3)} rotWithShape=1"


def test_a_turned_pictures_shadow_turns_with_its_rot(tmp_path: Path) -> None:
    shape = pictured(
        tmp_path,
        'width="10" height="10" filter="url(#f)" transform="rotate(90)"',
        defs='<filter id="f"><feDropShadow dx="3" dy="0" stdDeviation="1"/></filter>',
    )
    assert shape.find(f".//{A}outerShdw").get("dir") == "0"


# ---------------------------------------------------------------------------------------------
# clips
# ---------------------------------------------------------------------------------------------


def test_a_circle_clip_crops_to_its_box_and_draws_an_ellipse(tmp_path: Path) -> None:
    shape = pictured(
        tmp_path,
        'width="100" height="100" clip-path="url(#c)"',
        blob=SQUARE,
        defs='<clipPath id="c"><circle cx="50" cy="50" r="25"/></clipPath>',
    )
    assert box(shape) == (u(25), u(25), u(50), u(50))
    assert crop(shape) == {"b": "25000", "l": "25000", "r": "25000", "t": "25000"}
    assert geometry(shape) == ["prstGeom prst=ellipse", "avLst"]


@pytest.mark.parametrize(
    ("clip", "expected", "cropped", "drawn"),
    [
        (
            '<clipPath id="c"><ellipse cx="50" cy="50" rx="40" ry="20"/></clipPath>',
            (u(10), u(30), u(80), u(40)),
            {"b": "30000", "l": "10000", "r": "10000", "t": "30000"},
            ["prstGeom prst=ellipse", "avLst"],
        ),
        (
            '<clipPath id="c"><rect x="10" y="20" width="50" height="50"/></clipPath>',
            (u(10), u(20), u(50), u(50)),
            {"b": "30000", "l": "10000", "r": "40000", "t": "20000"},
            ["prstGeom prst=rect", "avLst"],
        ),
        (
            '<clipPath id="c" clipPathUnits="objectBoundingBox">'
            '<rect x="0.1" y="0.2" width="0.5" height="0.5" rx="0.1"/></clipPath>',
            (u(10), u(20), u(50), u(50)),
            {"b": "30000", "l": "10000", "r": "40000", "t": "20000"},
            ["prstGeom prst=roundRect", "avLst", "gd fmla=val 20000 name=adj"],
        ),
        (
            '<clipPath id="c" clipPathUnits="objectBoundingBox">'
            '<circle cx="0.5" cy="0.5" r="0.25"/></clipPath>',
            (u(25), u(25), u(50), u(50)),
            {"b": "25000", "l": "25000", "r": "25000", "t": "25000"},
            ["prstGeom prst=ellipse", "avLst"],
        ),
    ],
)
def test_a_clips_shape_is_the_pictures_geometry_in_its_box(
    tmp_path: Path,
    clip: str,
    expected: tuple[int, int, int, int],
    cropped: dict[str, str],
    drawn: list[str],
) -> None:
    shape = pictured(
        tmp_path, 'width="100" height="100" clip-path="url(#c)"', blob=SQUARE, defs=clip
    )
    assert box(shape) == expected
    assert crop(shape) == cropped
    assert geometry(shape) == drawn


@pytest.mark.parametrize(
    "clip",
    [
        '<clipPath id="c"><polygon points="50,0 100,100 0,100"/></clipPath>',
        '<clipPath id="c"><path d="M 50 0 L 100 100 L 0 100 Z"/></clipPath>',
        '<clipPath id="c" clipPathUnits="objectBoundingBox">'
        '<polygon points="0.5,0 1,1 0,1"/></clipPath>',
    ],
)
def test_a_polygon_or_path_clip_is_a_freeform_in_the_pictures_box(
    tmp_path: Path, clip: str
) -> None:
    shape = pictured(
        tmp_path, 'width="100" height="100" clip-path="url(#c)"', blob=SQUARE, defs=clip
    )
    assert box(shape) == (0, 0, u(100), u(100))
    assert crop(shape) == {}
    assert geometry(shape)[7:] == [
        f"path h={u(100)} w={u(100)}",
        "moveTo",
        f"pt x={u(50)} y=0",
        "lnTo",
        f"pt x={u(100)} y={u(100)}",
        "lnTo",
        f"pt x=0 y={u(100)}",
        "close",
    ]


def test_a_clip_past_the_shown_image_pads_it(tmp_path: Path) -> None:
    """`meet` shows a 200x100 picture at half size: a clip over the viewport runs past it."""
    shape = pictured(
        tmp_path,
        'width="100" height="100" clip-path="url(#c)"',
        defs='<clipPath id="c"><rect width="100" height="100"/></clipPath>',
    )
    assert box(shape) == (0, 0, u(100), u(100))
    assert crop(shape) == {"b": "-50000", "t": "-50000"}


def test_a_bounding_box_clip_reads_each_axis_by_its_own_side(tmp_path: Path) -> None:
    """A square picture in a 200x100 viewport shows at (50, 0, 100, 100); the clip's circle
    reaches a quarter of each side in from the viewport's."""
    shape = pictured(
        tmp_path,
        'width="200" height="100" clip-path="url(#c)"',
        blob=SQUARE,
        defs='<clipPath id="c" clipPathUnits="objectBoundingBox">'
        '<circle cx="0.5" cy="0.5" r="0.25"/></clipPath>',
    )
    assert box(shape) == (u(50), u(25), u(100), u(50))
    assert crop(shape) == {"b": "25000", "t": "25000"}


def test_a_clip_of_none_draws_the_picture_unclipped(tmp_path: Path) -> None:
    shape = pictured(tmp_path, 'width="100" height="100" clip-path="none"', blob=SQUARE)
    assert (box(shape), crop(shape)) == ((0, 0, u(100), u(100)), {})


def test_a_picture_with_no_width_crops_nothing(tmp_path: Path) -> None:
    shape = pictured(tmp_path, 'width="0" height="100"', blob=SQUARE)
    assert (box(shape)[2], crop(shape)) == (0, {})


def test_an_empty_clip_draws_the_picture_unclipped(tmp_path: Path) -> None:
    shape = pictured(
        tmp_path,
        'width="100" height="100" clip-path="url(#c)"',
        blob=SQUARE,
        defs='<clipPath id="c"/>',
    )
    assert (box(shape), crop(shape)) == ((0, 0, u(100), u(100)), {})
    assert geometry(shape) == ["prstGeom prst=rect", "avLst"]


# ---------------------------------------------------------------------------------------------
# preflight check 2, and refusals
# ---------------------------------------------------------------------------------------------


def test_a_slot_the_ledger_lacks_is_unfrozen(tmp_path: Path) -> None:
    with pytest.raises(TargetError) as caught:
        lowered(tmp_path, slide(f'<image href="{picture_slot(SQUARE)}"/>'))
    assert caught.value.code() == "OW_TARGET_ASSET_UNFROZEN"


@pytest.mark.parametrize(
    "blob",
    [
        b"<svg xmlns='http://www.w3.org/2000/svg'/>",
        b"not a picture at all",
        # a format PIL reads and PowerPoint does not hold
        b"P6\n1 1\n255\n\x00\x00\x00",
    ],
)
def test_an_asset_powerpoint_cannot_hold_is_refused_by_name(tmp_path: Path, blob: bytes) -> None:
    with pytest.raises(TargetError, match="is not a picture PowerPoint holds") as caught:
        lowered(tmp_path, slide(f'<image href="{picture_slot(blob)}"/>'), assets=[blob])
    assert caught.value.code() == "OW_TARGET_ASSET_FORMAT"


def test_a_clip_naming_no_clip_path_is_refused(tmp_path: Path) -> None:
    with pytest.raises(TargetError, match="it must name a <clipPath>") as caught:
        pictured(
            tmp_path,
            'width="10" height="10" clip-path="url(#f)"',
            defs='<filter id="f"><feGaussianBlur stdDeviation="1"/></filter>',
        )
    assert caught.value.code() == "OW_TARGET_REFERENCE_UNRESOLVED"
