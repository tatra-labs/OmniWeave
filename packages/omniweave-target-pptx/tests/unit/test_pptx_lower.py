"""`compile.pptx.native`'s `lower()` and `postflight()`: owsvg/1 to a `.pptx`. D795.

Every test builds a JobDir in `tmp_path` (`job.json`, `il/`, `assets/index.json`), lowers it with
`PptxCompiler` over an empty `FrozenLedger`, and reads the package back with `zipfile`, lxml and
python-pptx. The determinism tests compile twice in one process and once in a child with another
`PYTHONHASHSEED`, and compare part digests. D795 lowers the slide, groups, preset shapes, solid
paint, lines and transforms, and the preflight tests pin the refusal of what is not lowered yet.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import json
import os
import re
import subprocess  # noqa: TID251 - a second interpreter is the only witness for PYTHONHASHSEED.
import sys
from collections.abc import Sequence
from dataclasses import replace
from pathlib import Path

import pytest
from lowerkit import (  # type: ignore[import-not-found]  # tests/lowerkit.py
    CLOCK_NS,
    INK,
    RECT,
    TABLE,
    A,
    P,
    box,
    job_dir,
    line_of,
    lowered,
    outline,
    slide,
    turn,
    u,
)
from lxml import etree  # type: ignore[attr-defined] -- lxml ships no stubs
from omniweave_conform.harness import make_io
from omniweave_core.errors import TargetError
from omniweave_core.out.compile import lower_job
from omniweave_core.out.job import ILBundle, Job
from omniweave_core.out.lowering import resolve
from omniweave_core.out.opc import part_digest, write_opc
from omniweave_core.out.plan import UnitOutcome
from omniweave_core.out.receipt import Receipt
from omniweave_ports import ArtifactKind
from omniweave_target_pptx.driver import NORMALIZERS, PptxCompiler
from omniweave_target_pptx.lower import _pack
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

# ---------------------------------------------------------------------------------------------
# OUT1: the table's maps_to and lower/'s public functions, for the families D795 lowers
# ---------------------------------------------------------------------------------------------

NAMED = frozenset(row.maps_to for row in TABLE.rows if row.maps_to)
BUILT = frozenset(
    {
        "lower.defs:definition_id",
        "lower.defs:definitions",
        "lower.defs:reference",
        "lower.effects:apply",
        "lower.effects:crop",
        "lower.effects:crop_picture",
        "lower.effects:crop_units",
        "lower.effects:effect",
        "lower.effects:end_anchor",
        "lower.effects:end_size",
        "lower.effects:end_units",
        "lower.effects:glow",
        "lower.effects:head",
        "lower.effects:line_end",
        "lower.effects:orient",
        "lower.effects:outer_shadow",
        "lower.effects:radius",
        "lower.effects:shadow_alpha",
        "lower.effects:shadow_colour",
        "lower.effects:tail",
        "lower.effects:units",
        "lower.geometry:extent",
        "lower.geometry:offset",
        "lower.gradients:focus",
        "lower.gradients:linear",
        "lower.gradients:position",
        "lower.gradients:radial",
        "lower.gradients:stop",
        "lower.gradients:stop_alpha",
        "lower.gradients:stop_colour",
        "lower.gradients:units",
        "lower.group:group",
        "lower.ids:shape_name",
        "lower.media:aspect",
        "lower.media:picture",
        "lower.paint:alpha",
        "lower.paint:fill",
        "lower.paint:fill_alpha",
        "lower.paint:line",
        "lower.paint:line_alpha",
        "lower.paths:commands",
        "lower.paths:freeform",
        "lower.patterns:background",
        "lower.patterns:foreground",
        "lower.patterns:pattern",
        "lower.patterns:preset",
        "lower.shapes:centre",
        "lower.shapes:corner",
        "lower.shapes:ellipse",
        "lower.shapes:endpoint",
        "lower.shapes:freeform",
        "lower.shapes:line",
        "lower.shapes:radius",
        "lower.shapes:rect",
        "lower.shapes:vertices",
        "lower.slide:size",
        "lower.slide:slide",
        "lower.stroke:cap",
        "lower.stroke:dash",
        "lower.stroke:join",
        "lower.stroke:scaling",
        "lower.stroke:width",
        "lower.text:anchor",
        "lower.text:baseline",
        "lower.text:decoration",
        "lower.text:italic",
        "lower.text:run",
        "lower.text:shift",
        "lower.text:size",
        "lower.text:spacing",
        "lower.text:text_box",
        "lower.text:typeface",
        "lower.text:weight",
        "lower.text:whitespace",
        "lower.transform:transform",
    }
)


def test_the_rows_lowered_so_far_resolve_and_no_others_do() -> None:
    """D795's 26 functions, then D796's freeforms, D797's text, D798's paint servers and
    D799's effects, and D800's pictures."""
    assert {m for m in NAMED if resolve(m, _pack.FAMILIES) is not None} == BUILT


def test_every_public_lowering_function_is_one_a_row_names() -> None:
    """OUT1's other direction: `lower/` has no function the table does not declare."""
    for family, module in _pack.FAMILIES.items():
        exported = {f"{family}:{n}" for n in module.__all__ if callable(getattr(module, n))}
        assert exported, family
        assert exported <= NAMED, (family, sorted(exported - NAMED))


# ---------------------------------------------------------------------------------------------
# preflight (09 section 5.1)
# ---------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("body", "named"),
    [
        (
            '<defs><symbol id="s"><rect width="1" height="1"/></symbol></defs>',
            "def.symbol (lower.defs:symbol)",
        ),
        (
            '<g data-ow-manual="unverifiable_data"><rect width="1" height="1"/></g>',
            "ow.manual (lower.manual:placeholder)",
        ),
        ('<a href="#u02"><rect width="1" height="1"/></a>', "shape.a (lower.links:hyperlink)"),
        ('<use href="#r"/>', "def.use (lower.defs:use)"),
        (
            '<rect width="1" height="1" data-ow-cite="c1"/>',
            "ow.cite (lower.cites:provenance_part)",
        ),
    ],
)
def test_a_construct_not_yet_lowered_is_refused_before_a_byte(
    tmp_path: Path, body: str, named: str
) -> None:
    with pytest.raises(TargetError, match=re.escape(named)) as caught:
        lowered(tmp_path, slide(f"{RECT}{body}"))
    assert caught.value.code() == "OW_TARGET_CONSTRUCT_UNMAPPED"
    assert "shape.rect" not in str(caught.value)
    assert (caught.value.fix or "").startswith("ow out constructs pptx --construct ")
    assert not (tmp_path / "deck.pptx").exists()


@pytest.mark.parametrize(
    ("attrs", "row"),
    [
        ('fill="#1f4e79"', "paint.fill.hex"),
        ('x="10px"', "geom.x.px"),
        ('opacity="50%"', "paint.opacity.percent"),
        ('stroke-width="2px"', "stroke.width.px"),
    ],
)
def test_a_value_a_normalize_row_would_rewrite_is_refused(
    tmp_path: Path, attrs: str, row: str
) -> None:
    with pytest.raises(
        TargetError, match=f"1 value\\(s\\) are not in canonical form: {row} at "
    ) as caught:
        lowered(tmp_path, slide(f'<rect width="1" height="1" {attrs}/>'))
    assert caught.value.code() == "OW_TARGET_NORMALIZE_AT_COMPILE"


def test_the_normalize_refusal_names_three_and_counts_them_all(tmp_path: Path) -> None:
    with pytest.raises(TargetError, match="4 value") as caught:
        lowered(tmp_path, slide('<rect x="1px" y="1px" width="1px" height="1px"/>'))
    assert str(caught.value).count(" at ") == 3


def test_the_plans_roster_is_the_bundles(tmp_path: Path) -> None:
    with pytest.raises(TargetError, match="are not the bundle's") as caught:
        lowered(tmp_path, slide(RECT), slide(RECT), plan_units=["u01"])
    assert caught.value.code() == "OW_TARGET_ROSTER_MISMATCH"


def test_the_cards_max_units_is_a_ceiling(tmp_path: Path) -> None:
    job, plan, ledger, _ = job_dir(tmp_path, [slide(RECT), slide(RECT)])
    compiler = PptxCompiler()
    bundle = ILBundle.load(job / "il", carrier=compiler.carrier, constructs=TABLE)

    def build(max_units: int) -> bytes:
        return _pack.build(
            plan,
            bundle,
            ledger,
            carrier=compiler.carrier,
            table=TABLE,
            normalizers=NORMALIZERS,
            clock_ns=CLOCK_NS,
            max_units=max_units,
        )

    assert build(2).startswith(b"PK")
    with pytest.raises(TargetError, match="max_units is 1") as caught:
        build(1)
    assert caught.value.code() == "OW_TARGET_TOO_MANY_UNITS"


@pytest.mark.parametrize(
    ("boxes", "named"),
    [
        (("0 0 1280 720", "0 0 1920 1080"), "1280x720, 1920x1080"),
        (("0 0 1280 720", "0 0 1280 720", "0 0 960 540"), "1280x720, 1280x720, 960x540"),
        (("10 10 1280 720",), "none"),
    ],
)
def test_a_deck_has_one_slide_size(tmp_path: Path, boxes: Sequence[str], named: str) -> None:
    with pytest.raises(TargetError, match=f"its units' viewBoxes are {named}") as caught:
        lowered(tmp_path, *(slide(RECT, box=b) for b in boxes))
    assert caught.value.code() == "OW_TARGET_CANVAS_MISMATCH"


def test_an_element_no_native_row_declares_is_refused_by_name(tmp_path: Path) -> None:
    """The gate refuses `<foreignObject>`; a lowering handed one names it, not skips it."""
    with pytest.raises(TargetError, match=r"<foreignObject> at .* matches no native row") as caught:
        lowered(tmp_path, slide("<foreignObject/>"))
    assert caught.value.code() == "OW_TARGET_CONSTRUCT_UNMAPPED"


@pytest.mark.parametrize(
    ("body", "said"),
    [
        ('<rect width="1" height="1" fill="url(#nowhere)"/>', r"url\(#nowhere\) names no id"),
        (
            '<rect id="r" width="1" height="1"/><rect width="1" height="1" stroke="url(#r)"/>',
            r"url\(#r\) names a <rect>",
        ),
        ('<rect width="1" height="1" fill="url(nowhere)"/>', "names no id"),
    ],
)
def test_a_paint_that_names_no_paint_is_refused(tmp_path: Path, body: str, said: str) -> None:
    with pytest.raises(TargetError, match=said) as caught:
        lowered(tmp_path, slide(body))
    assert caught.value.code() == "OW_TARGET_REFERENCE_UNRESOLVED"


# ---------------------------------------------------------------------------------------------
# the package (09 sections 6 and 12.1)
# ---------------------------------------------------------------------------------------------


def test_the_package_holds_a_deck_and_none_of_the_templates_baggage(tmp_path: Path) -> None:
    out = lowered(tmp_path, slide(RECT), slide(RECT))
    names = set(out.members)
    assert {"ppt/presentation.xml", "ppt/slides/slide1.xml", "ppt/slides/slide2.xml"} <= names
    assert not [n for n in names if "thumbnail" in n or "printerSettings" in n]
    assert not [n for n, data in out.members.items() if b"thumbnail" in data]
    assert list(out.members)[:2] == ["[Content_Types].xml", "_rels/.rels"]


def test_every_attribute_map_is_sorted(tmp_path: Path) -> None:
    out = lowered(tmp_path, slide(f'<rect x="1" y="2" width="3" height="4" {INK}/>'))
    for name, payload in out.members.items():
        if not name.endswith((".xml", ".rels")):
            continue
        for element in etree.fromstring(payload).iter():
            if isinstance(element.tag, str):
                keys = list(element.attrib.keys())
                assert keys == sorted(keys), (name, element.tag)


def test_the_core_properties_are_the_jobs_clock_and_no_host(tmp_path: Path) -> None:
    out = lowered(tmp_path, slide(RECT))
    core = out.deck().core_properties
    when = dt.datetime.fromtimestamp(CLOCK_NS // 1_000_000_000, tz=dt.UTC).replace(tzinfo=None)
    assert core.created.replace(tzinfo=None) == when
    assert core.modified.replace(tzinfo=None) == when
    assert (core.revision, core.author, core.last_modified_by, core.title, core.comments) == (
        1,
        "",
        "",
        "deck",
        "",
    )
    app = out.xml("docProps/app.xml")
    assert [(etree.QName(c).localname, c.text) for c in app] == [
        ("Application", "omniweave-target-pptx"),
        ("DocSecurity", "0"),
    ]


@pytest.mark.parametrize(
    ("box_", "size"),
    [("0 0 1280 720", (12192000, 6858000)), ("0 0 960 540", (9144000, 5143500))],
)
def test_the_slide_size_is_the_viewbox_and_names_no_template_type(
    tmp_path: Path, box_: str, size: tuple[int, int]
) -> None:
    out = lowered(tmp_path, slide(RECT, box=box_))
    deck = out.deck()
    assert (deck.slide_width, deck.slide_height) == size
    [sized] = out.xml("ppt/presentation.xml").iter(f"{P}sldSz")
    assert dict(sized.attrib) == {"cx": str(size[0]), "cy": str(size[1])}


def test_each_unit_is_a_slide_in_plan_order_with_ids_from_two(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide(f'<rect data-ow-el="u01/e2" width="1" height="1"/>{RECT}<circle r="1"/>'),
        slide('<rect data-ow-el="u02/e2" width="1" height="1"/>'),
    )
    assert [[(s.shape_id, s.name) for s in page.shapes] for page in out.deck().slides] == [
        [(2, "u01/e2"), (3, "rect 3"), (4, "circle 4")],
        [(2, "u02/e2")],
    ]


def test_what_the_table_ignores_and_a_defs_draw_nothing(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide(
            f"<title>t</title><desc>d</desc><metadata/><defs/>{RECT}"
            '<rect id="b" width="1" height="1"/>'
        ),
    )
    assert [s.find(f"{P}nvSpPr/{P}cNvPr").get("name") for s in out.shapes()] == [
        "rect 2",
        "rect 3",
    ]


def test_two_compiles_are_one_package_byte_for_byte(tmp_path: Path) -> None:
    units = (
        slide(f'<g transform="rotate(30)"><rect width="10" height="5" {INK}/></g>'),
        slide('<ellipse cx="5" cy="5" rx="5" ry="2" fill-opacity="0.5"/>'),
    )
    first = lowered(tmp_path / "one", *units)
    second = lowered(tmp_path / "two", *units)
    assert first.data == second.data


_CHILD = """
import sys
sys.path.insert(0, sys.argv[1])
from lowerkit import lowered, slide
from pathlib import Path
from omniweave_core.out.opc import part_digest
out = lowered(Path(sys.argv[2]), slide('<g opacity="0.5"><rect width="9" height="4"/></g>'))
print(part_digest(out.members))
"""


def test_the_part_digest_holds_under_another_hash_seed(tmp_path: Path) -> None:
    """09 section 6: `part` is promised; a child with `PYTHONHASHSEED=1` lowers the same parts."""
    here = lowered(tmp_path / "here", slide('<g opacity="0.5"><rect width="9" height="4"/></g>'))
    env = {**os.environ, "PYTHONHASHSEED": "1"}
    child = subprocess.run(  # noqa: S603 -- this interpreter, this test's own module
        [sys.executable, "-c", _CHILD, str(Path(__file__).parents[1]), str(tmp_path / "there")],
        capture_output=True,
        text=True,
        env=env,
        check=True,
        timeout=120,
    )
    assert child.stdout.strip() == part_digest(here.members)


# ---------------------------------------------------------------------------------------------
# postflight (09 section 5.3)
# ---------------------------------------------------------------------------------------------


def test_postflight_checks_the_package_and_reports_the_classification(tmp_path: Path) -> None:
    out = lowered(tmp_path, slide(RECT), slide(RECT))
    receipt = out.receipt()
    assert (receipt.target, receipt.slug, receipt.artifact_gen) == ("pptx", "deck", 1)
    assert receipt.part_digest == part_digest(out.members).removeprefix("sha256:")
    assert receipt.output == {
        "path": str(out.path),
        "bytes": len(out.data),
        "sha256": hashlib.sha256(out.data).hexdigest(),
    }
    assert dict(receipt.units) == {"u01": UnitOutcome.NATIVE, "u02": UnitOutcome.NATIVE}
    assert dict(receipt.counts) == {UnitOutcome.NATIVE: 2}
    checks = {c.name: (c.value, c.detail) for c in receipt.checks}
    assert checks["slide_count"] == ("passed", "2 slides == unit_count")
    assert checks["slide_size"] == ("passed", "12192000x6858000 EMU")
    assert checks["citations_channel"][0] == "skipped"
    assert (checks["bake_used"][0], checks["measure_available"][0]) == (
        "not-applicable",
        "not-applicable",
    )
    assert {checks[n][0] for n in ("zip_integrity", "content_types")} == {"passed"}
    [enforced] = [c for c in receipt.checks if c.name == "internal_relationships"]
    assert enforced.value == "enforced-at-build"
    path, _, name = (enforced.fixture or "").partition("::")
    assert Path(__file__).as_posix().endswith(path)
    assert name == test_a_dangling_relationship_is_refused_at_build.__name__
    assert json.loads(receipt.to_bytes())["target"] == "pptx"


def test_a_dangling_relationship_is_refused_at_build(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """09 section 5.3: `enforced-at-build` is proved by breaking the part graph."""
    canonical = _pack._canonical

    def broken(name: str, data: bytes) -> bytes:
        made = canonical(name, data)
        if name == "ppt/_rels/presentation.xml.rels":
            made = made.replace(
                b"</Relationships>",
                b'<Relationship Id="rId99" Target="slides/slide9.xml" Type="x"/></Relationships>',
            )
        return made

    monkeypatch.setattr(_pack, "_canonical", broken)
    with pytest.raises(TargetError, match=r"slide9\.xml") as caught:
        lowered(tmp_path, slide(RECT))
    assert caught.value.code() == "OW_TARGET_DANGLING_REL"


def test_a_package_with_a_unit_missing_is_refused(tmp_path: Path) -> None:
    out = lowered(tmp_path, slide(RECT), slide(RECT))
    third = replace(out.plan.units[-1], id="u03", ord=3, il_path="il/u03.svg")
    plan = replace(out.plan, unit_count=3, units=(*out.plan.units, third))
    with pytest.raises(TargetError, match="holds 2 units and the plan commissioned 3") as caught:
        PptxCompiler().postflight(str(out.path), plan)
    assert caught.value.code() == "OW_TARGET_UNIT_COUNT_MISMATCH"


def test_a_presentation_with_no_slide_size_says_none(tmp_path: Path) -> None:
    out = lowered(tmp_path, slide(RECT))
    deck = out.xml("ppt/presentation.xml")
    [sized] = deck.iter(f"{P}sldSz")
    deck.remove(sized)
    members = {**out.members, "ppt/presentation.xml": etree.tostring(deck)}
    out.path.write_bytes(write_opc(members))
    [size] = [c for c in out.receipt().checks if c.name == "slide_size"]
    assert size.detail == "none"


def test_lower_job_runs_the_worker_half_over_a_pptxjob_dir(tmp_path: Path) -> None:
    """D771's `lower_job()`: the JobDir's digests held, lowered, staged, post-flighted."""
    job, plan, _, store = job_dir(tmp_path, [slide(RECT)])
    result = lower_job(
        PptxCompiler(),
        Job.load(job),
        job,
        constructs=TABLE,
        suffix=".pptx",
        io=make_io(store, tmp_path / "io"),
    )
    artifact, receipt = result.produced
    assert (artifact.kind, receipt.kind) == (ArtifactKind.ARTIFACT, ArtifactKind.RECEIPT)
    body = receipt.inline if receipt.inline is not None else store.open(receipt.blob or "").read()
    assert Receipt.from_json(json.loads(body)).slug == plan.slug


# ---------------------------------------------------------------------------------------------
# shapes
# ---------------------------------------------------------------------------------------------


def test_a_rect_is_a_preset_shape_in_drawingmls_order(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide('<rect data-ow-el="u01/e2" x="10" y="20" width="100" height="50" fill="#1F4E79"/>'),
    )
    [shape] = out.shapes()
    assert outline(shape) == [
        "sp",
        "nvSpPr",
        "cNvPr id=2 name=u01/e2",
        "cNvSpPr",
        "nvPr",
        "spPr",
        "xfrm",
        "off x=95250 y=190500",
        "ext cx=952500 cy=476250",
        "prstGeom prst=rect",
        "avLst",
        "solidFill",
        "srgbClr val=1F4E79",
        "ln cap=flat w=0",
        "noFill",
        "miter lim=800000",
    ]


@pytest.mark.parametrize(
    ("attrs", "adj"),
    [
        ('width="100" height="50" rx="10"', 20000),
        ('width="50" height="100" rx="10"', 20000),
        ('width="100" height="50" ry="5"', 10000),
        ('width="100" height="50" rx="4" ry="10"', 20000),
        ('width="100" height="50" rx="10" ry="4"', 20000),
        ('width="100" height="50" rx="40"', 50000),
        ('width="100" height="50" rx="0"', None),
        ('width="0" height="50" rx="10"', None),
    ],
)
def test_a_rounded_rect_is_round_rect_by_its_larger_radius(
    tmp_path: Path, attrs: str, adj: int | None
) -> None:
    out = lowered(tmp_path, slide(f"<rect {attrs}/>"))
    geometry = out.first().find(f"{P}spPr/{A}prstGeom")
    if adj is None:
        assert outline(geometry) == ["prstGeom prst=rect", "avLst"]
    else:
        assert outline(geometry) == [
            "prstGeom prst=roundRect",
            "avLst",
            f"gd fmla=val {adj} name=adj",
        ]


@pytest.mark.parametrize(
    ("body", "expected"),
    [
        ('<circle cx="100" cy="100" r="50"/>', (u(50), u(50), u(100), u(100))),
        ('<ellipse cx="100" cy="50" rx="60" ry="30"/>', (u(40), u(20), u(120), u(60))),
        ('<ellipse rx="6" ry="3" r="9"/>', (u(-6), u(-3), u(12), u(6))),
    ],
)
def test_circles_and_ellipses_are_the_ellipse_preset_over_their_box(
    tmp_path: Path, body: str, expected: tuple[int, int, int, int]
) -> None:
    shape = lowered(tmp_path, slide(body)).first()
    assert box(shape) == expected
    assert shape.find(f"{P}spPr/{A}prstGeom").get("prst") == "ellipse"


@pytest.mark.parametrize(
    ("ends", "expected", "flips"),
    [
        ('x1="10" y1="10" x2="100" y2="40"', (u(10), u(10), u(90), u(30)), (None, None)),
        ('x1="100" y1="100" x2="10" y2="10"', (u(10), u(10), u(90), u(90)), ("1", "1")),
        ('x1="10" y1="40" x2="100" y2="10"', (u(10), u(10), u(90), u(30)), (None, "1")),
        ('x1="100" y1="10" x2="10" y2="40"', (u(10), u(10), u(90), u(30)), ("1", None)),
        ('x2="10"', (0, 0, u(10), 0), (None, None)),
    ],
)
def test_a_line_spans_its_ends_and_its_flips_say_which_way_it_runs(
    tmp_path: Path,
    ends: str,
    expected: tuple[int, int, int, int],
    flips: tuple[str | None, str | None],
) -> None:
    shape = lowered(tmp_path, slide(f'<line {ends} {INK} fill="#FF0000"/>')).first()
    assert box(shape) == expected
    assert turn(shape) == (None, *flips)
    body = shape.find(f"{P}spPr")
    assert [etree.QName(c).localname for c in body] == ["xfrm", "prstGeom", "noFill", "ln"]
    assert body.find(f"{A}prstGeom").get("prst") == "line"


def test_a_line_through_a_transform_takes_the_matrix_at_its_ends(tmp_path: Path) -> None:
    shape = lowered(
        tmp_path, slide(f'<g transform="translate(5 5) scale(2)"><line x2="10" {INK}/></g>')
    ).first()
    assert box(shape) == (u(5), u(5), u(20), 0)


def test_python_pptx_reads_the_shapes_back(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide(
            '<rect data-ow-el="u01/e2" x="10" y="20" width="100" height="50"/>'
            '<g><ellipse cx="5" cy="5" rx="5" ry="5"/></g>'
        ),
    )
    [page] = out.deck().slides
    rect, group = page.shapes
    assert (rect.name, rect.shape_id, rect.auto_shape_type) == ("u01/e2", 2, MSO_SHAPE.RECTANGLE)
    assert (rect.left, rect.top, rect.width, rect.height) == (u(10), u(20), u(100), u(50))
    assert group.shape_type == MSO_SHAPE_TYPE.GROUP
    [inner] = group.shapes
    assert (inner.auto_shape_type, inner.shape_id) == (MSO_SHAPE.OVAL, 4)


# ---------------------------------------------------------------------------------------------
# groups
# ---------------------------------------------------------------------------------------------


def test_a_group_frames_its_children_one_to_one(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide(
            '<g data-ow-el="u01/e2"><rect x="10" y="10" width="10" height="10"/>'
            '<rect x="50" y="30" width="20" height="20"/></g>'
        ),
    )
    [made] = out.shapes()
    assert outline(made)[:12] == [
        "grpSp",
        "nvGrpSpPr",
        "cNvPr id=2 name=u01/e2",
        "cNvGrpSpPr",
        "nvPr",
        "grpSpPr",
        "xfrm",
        f"off x={u(10)} y={u(10)}",
        f"ext cx={u(60)} cy={u(40)}",
        f"chOff x={u(10)} y={u(10)}",
        f"chExt cx={u(60)} cy={u(40)}",
        "sp",
    ]
    children = made.findall(f"{P}sp")
    assert [box(c) for c in children] == [
        (u(10), u(10), u(10), u(10)),
        (u(50), u(30), u(20), u(20)),
    ]
    assert [c.find(f"{P}nvSpPr/{P}cNvPr").get("id") for c in children] == ["3", "4"]


def test_an_empty_group_has_no_extent_and_a_nested_one_counts_in_its_parents_box(
    tmp_path: Path,
) -> None:
    out = lowered(
        tmp_path,
        slide(
            '<g/><g><g><rect x="5" y="5" width="5" height="5"/></g>'
            '<rect x="20" y="20" width="10" height="10"/></g>'
        ),
    )
    empty, outer = out.shapes()
    assert box(empty) == (0, 0, 0, 0)
    assert box(outer) == (u(5), u(5), u(25), u(25))
    [inner] = outer.findall(f"{P}grpSp")
    assert box(inner) == (u(5), u(5), u(5), u(5))
    assert [c.get("id") for c in out.tree().iter(f"{P}cNvPr")] == ["1", "2", "3", "4", "5", "6"]


# ---------------------------------------------------------------------------------------------
# transforms
# ---------------------------------------------------------------------------------------------

UNTURNED = (None, None, None)


@pytest.mark.parametrize(
    ("body", "expected", "turned"),
    [
        (
            '<g transform="translate(100 50)"><rect width="10" height="10"/></g>',
            (u(100), u(50), u(10), u(10)),
            UNTURNED,
        ),
        (
            '<g transform="translate(100)"><rect width="10" height="10"/></g>',
            (u(100), 0, u(10), u(10)),
            UNTURNED,
        ),
        (
            '<rect x="10" y="10" width="10" height="10" transform="scale(2)"/>',
            (u(20), u(20), u(20), u(20)),
            UNTURNED,
        ),
        (
            '<rect x="10" y="10" width="10" height="10" transform="scale(2 3)"/>',
            (u(20), u(30), u(20), u(30)),
            UNTURNED,
        ),
        (
            '<rect width="20" height="10" transform="rotate(90 10 5)"/>',
            (0, 0, u(20), u(10)),
            ("5400000", None, None),
        ),
        (
            '<rect width="20" height="10" transform="rotate(-90 10 5)"/>',
            (0, 0, u(20), u(10)),
            ("16200000", None, None),
        ),
        (
            '<rect x="10" width="10" height="10" transform="translate(100 0) scale(-1 1)"/>',
            (u(80), 0, u(10), u(10)),
            ("10800000", None, "1"),
        ),
        (
            '<g transform="translate(10 0)"><g transform="rotate(90)">'
            '<rect width="10" height="10"/></g></g>',
            (0, 0, u(10), u(10)),
            ("5400000", None, None),
        ),
        (
            '<rect width="10" height="10" transform="matrix(1 0 0 1 5 6)"/>',
            (u(5), u(6), u(10), u(10)),
            UNTURNED,
        ),
    ],
)
def test_a_preset_shape_takes_its_matrix_as_a_box_a_turn_and_a_flip(
    tmp_path: Path,
    body: str,
    expected: tuple[int, int, int, int],
    turned: tuple[str | None, str | None, str | None],
) -> None:
    shape = lowered(tmp_path, slide(body)).first()
    assert box(shape) == expected
    assert turn(shape) == turned


def test_the_slide_roots_own_transform_is_left_out(tmp_path: Path) -> None:
    text = (
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720" transform="translate(5 5)">'
        '<rect x="1" y="1" width="1" height="1"/></svg>'
    )
    assert box(lowered(tmp_path, text).first()) == (u(1), u(1), u(1), u(1))


# ---------------------------------------------------------------------------------------------
# paint
# ---------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("body", "fill"),
    [
        (RECT, ["solidFill", "srgbClr val=000000"]),
        ('<rect width="1" height="1" fill="none"/>', ["noFill"]),
        ('<g fill="#00FF00"><rect width="1" height="1"/></g>', ["solidFill", "srgbClr val=00FF00"]),
        (
            '<rect width="1" height="1" fill="#FF0000" style="fill: #0000FF"/>',
            ["solidFill", "srgbClr val=0000FF"],
        ),
        (
            '<g style="fill: #00FF00"><rect width="1" height="1" fill="#FF0000"/></g>',
            ["solidFill", "srgbClr val=FF0000"],
        ),
        # no normalize row settles a `style` value, so the colour is settled where it is written
        ('<rect width="1" height="1" style="fill: #f00"/>', ["solidFill", "srgbClr val=FF0000"]),
    ],
)
def test_a_fill_is_inherited_and_style_outranks_the_attribute(
    tmp_path: Path, body: str, fill: list[str]
) -> None:
    shape = lowered(tmp_path, slide(body)).first()
    assert outline(shape.find(f"{P}spPr")[2]) == fill


@pytest.mark.parametrize(
    ("body", "level"),
    [
        (RECT, None),
        ('<rect width="1" height="1" fill-opacity="0.25"/>', "25000"),
        ('<g opacity="0.5"><rect width="1" height="1" opacity="0.5"/></g>', "25000"),
        ('<g fill-opacity="0.5"><rect width="1" height="1"/></g>', "50000"),
        (
            '<g fill-opacity="0.5" opacity="0.5">'
            '<rect width="1" height="1" fill-opacity="0.4"/></g>',
            "20000",
        ),
        ('<rect width="1" height="1" opacity="2" fill-opacity="0.5"/>', "50000"),
        ('<rect width="1" height="1" opacity="-1"/>', "0"),
        ('<rect width="1" height="1" fill-opacity="1.5" opacity="0.5"/>', "50000"),
        ('<rect width="1" height="1" style="opacity: 0.5" opacity="0.2"/>', "50000"),
    ],
)
def test_alpha_is_fill_opacity_inherited_times_opacity_down_the_chain(
    tmp_path: Path, body: str, level: str | None
) -> None:
    [colour] = lowered(tmp_path, slide(body)).tree().iter(f"{A}srgbClr")
    alpha = colour.find(f"{A}alpha")
    assert (None if alpha is None else alpha.get("val")) == level


def test_a_stroke_is_the_lines_paint_and_none_draws_none(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide(
            '<rect width="1" height="1" stroke="#112233" stroke-opacity="0.5" opacity="0.5"/>'
            '<rect width="1" height="1" stroke="none" stroke-width="4"/>'
            '<g stroke="#445566"><rect width="1" height="1"/></g>'
        ),
    )
    first, second, inner = out.tree().iter(f"{P}sp")
    assert outline(line_of(first)) == [
        "ln cap=flat w=9525",
        "solidFill",
        "srgbClr val=112233",
        "alpha val=25000",
        "miter lim=800000",
    ]
    assert outline(line_of(second)) == ["ln cap=flat w=0", "noFill", "miter lim=800000"]
    assert outline(line_of(inner))[1:3] == ["solidFill", "srgbClr val=445566"]


# ---------------------------------------------------------------------------------------------
# stroke
# ---------------------------------------------------------------------------------------------


def _line(tmp_path: Path, attrs: str) -> list[str]:
    return outline(line_of(lowered(tmp_path, slide(f"<rect {INK} {attrs}/>")).first()))


@pytest.mark.parametrize(
    ("attrs", "cap"),
    [
        ("", "flat"),
        ('stroke-linecap="butt"', "flat"),
        ('stroke-linecap="round"', "rnd"),
        ('stroke-linecap="square"', "sq"),
    ],
)
def test_a_cap_is_drawingmls_word_for_it(tmp_path: Path, attrs: str, cap: str) -> None:
    assert _line(tmp_path, attrs)[0] == f"ln cap={cap} w=9525"


@pytest.mark.parametrize(
    ("attrs", "join"),
    [
        ("", "miter lim=800000"),
        ('stroke-linejoin="miter"', "miter lim=800000"),
        ('stroke-linejoin="round"', "round"),
        ('stroke-linejoin="bevel"', "bevel"),
    ],
)
def test_a_join_is_its_element(tmp_path: Path, attrs: str, join: str) -> None:
    assert _line(tmp_path, attrs)[-1] == join


@pytest.mark.parametrize(
    ("attrs", "dash"),
    [
        ("", []),
        ('stroke-dasharray="none"', []),
        ('stroke-dasharray="4 4"', ["prstDash val=dash"]),
        ('stroke-dasharray="6 3"', ["prstDash val=dash"]),
        ('stroke-dasharray="2 2"', ["prstDash val=sysDot"]),
        ('stroke-dasharray="8,4"', ["prstDash val=lgDash"]),
        ('stroke-dasharray="8 4 2 4"', ["prstDash val=lgDashDot"]),
        ('stroke-dasharray="10 5" stroke-width="2"', ["custDash", "ds d=500000 sp=250000"]),
        ('stroke-dasharray="10 5" stroke-width="0"', ["custDash", "ds d=1000000 sp=500000"]),
    ],
)
def test_a_dash_is_a_preset_or_its_first_dash_and_gap(
    tmp_path: Path, attrs: str, dash: list[str]
) -> None:
    assert _line(tmp_path, attrs)[3:-1] == dash


@pytest.mark.parametrize(
    ("body", "width"),
    [
        (f"<rect {INK}/>", u(1)),
        (f'<rect {INK} stroke-width="2"/>', u(2)),
        (f'<g stroke-width="3"><rect {INK}/></g>', u(3)),
        (f'<g transform="scale(2)"><rect {INK} stroke-width="3"/></g>', u(6)),
        (f'<g transform="scale(4 1)"><rect {INK} stroke-width="3"/></g>', u(6)),
        (
            f'<g transform="scale(2)"><rect {INK} stroke-width="3" '
            'vector-effect="non-scaling-stroke"/></g>',
            u(3),
        ),
        (
            f'<g transform="scale(2)" vector-effect="non-scaling-stroke">'
            f'<rect {INK} stroke-width="3"/></g>',
            u(6),
        ),
        ('<rect stroke-width="3"/>', 0),
    ],
)
def test_a_width_grows_with_the_scale_unless_the_node_says_it_does_not(
    tmp_path: Path, body: str, width: int
) -> None:
    assert line_of(lowered(tmp_path, slide(body)).first()).get("w") == str(width)
