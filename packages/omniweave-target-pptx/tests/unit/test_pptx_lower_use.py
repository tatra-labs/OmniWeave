"""`compile.pptx.native`'s `<use>`, `<symbol>` and links. D801.

A `<use>` is a group named by its address, holding a clone of what it names; a `<symbol>`'s
clone is its content, its `viewBox` fitted into the `<use>`'s box. An `<a>` puts a click on every
shape it holds: a page as an external hyperlink, `#u<NN>` as a jump to that unit's slide. Each is
lowered through `lowered()` (tests/lowerkit.py) and read back from the package and with
python-pptx.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from lowerkit import (  # type: ignore[import-not-found]  # tests/lowerkit.py
    A,
    P,
    box,
    lowered,
    picture_slot,
    png,
    slide,
    u,
)
from lxml import etree  # type: ignore[attr-defined] -- lxml ships no stubs
from omniweave_core.errors import TargetError
from pptx.oxml.ns import qn

RECT = '<rect id="r" x="1" y="2" width="10" height="10"/>'


def names(tree: Any) -> list[str]:
    return [c.get("name") for c in tree.iter(f"{P}cNvPr")][1:]


def first_sp(made: Any) -> Any:
    return next(made.iter(f"{P}sp"))


def clicks(tree: Any) -> list[tuple[str, str | None]]:
    """Each shape's name and its click's action, `None` for no click."""
    out = []
    for named in list(tree.iter(f"{P}cNvPr"))[1:]:
        click = named.find(f"{A}hlinkClick")
        out.append((named.get("name"), None if click is None else click.get("action", "")))
    return out


# ---------------------------------------------------------------------------------------------
# <use>
# ---------------------------------------------------------------------------------------------


def test_a_use_is_a_group_named_by_it_holding_a_clone(tmp_path: Path) -> None:
    out = lowered(
        tmp_path, slide(f'<defs>{RECT}</defs><use data-ow-el="u01/e2" href="#r" x="100" y="50"/>')
    )
    [made] = out.shapes()
    assert made.tag == f"{P}grpSp"
    assert names(out.tree()) == ["u01/e2", "rect 3"]
    assert box(first_sp(made)) == (u(101), u(52), u(10), u(10))
    assert box(made) == (u(101), u(52), u(10), u(10))


def test_a_clone_of_a_drawn_shape_takes_no_address(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide(
            '<rect id="r" data-ow-el="u01/e2" width="10" height="10"/>'
            '<use data-ow-el="u01/e3" href="#r" x="20"/>'
        ),
    )
    assert names(out.tree()) == ["u01/e2", "u01/e3", "rect 4"]


@pytest.mark.parametrize(
    ("body", "colour"),
    [
        (
            '<defs><rect id="r" width="1" height="1"/></defs><use href="#r" fill="#00FF00"/>',
            "00FF00",
        ),
        (
            '<defs><g fill="#0000FF"><rect id="r" width="1" height="1"/></g></defs>'
            '<use href="#r"/>',
            "000000",
        ),
        (
            '<defs><rect id="r" width="1" height="1" fill="#FF0000"/></defs>'
            '<use href="#r" fill="#00FF00"/>',
            "FF0000",
        ),
    ],
)
def test_a_clone_inherits_from_its_use_not_its_original_parent(
    tmp_path: Path, body: str, colour: str
) -> None:
    out = lowered(tmp_path, slide(body))
    assert first_sp(out.tree()).find(f"{P}spPr/{A}solidFill/{A}srgbClr").get("val") == colour


def test_a_uses_transform_comes_before_its_x_and_y(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide(f'<defs>{RECT}</defs><use href="#r" x="10" transform="scale(2)"/>'),
    )
    assert box(first_sp(out.tree())) == (u(22), u(4), u(20), u(20))


def test_a_use_of_a_use_nests_and_adds_its_moves(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide(f'<defs>{RECT}<use id="v" href="#r" x="5"/></defs><use href="#v" x="10"/>'),
    )
    [outer] = out.shapes()
    [inner] = outer.findall(f"{P}grpSp")
    assert box(first_sp(inner)) == (u(16), u(2), u(10), u(10))
    assert names(out.tree()) == ["use 2", "use 3", "rect 4"]


@pytest.mark.parametrize(
    ("target", "kind"),
    [
        ('<text id="t" y="20">a</text>', f"{P}sp"),
        (f'<image id="t" href="{picture_slot(png(2, 2))}" width="2" height="2"/>', f"{P}pic"),
        ('<g id="t"><rect width="1" height="1"/></g>', f"{P}grpSp"),
    ],
)
def test_a_use_clones_a_text_a_picture_or_a_group(tmp_path: Path, target: str, kind: str) -> None:
    out = lowered(tmp_path, slide(f'<defs>{target}</defs><use href="#t"/>'), assets=[png(2, 2)])
    [made] = out.shapes()
    assert [c.tag for c in made if c.tag.startswith(P) and not c.tag.endswith("Pr")] == [kind]


def test_a_uses_filter_is_its_groups(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide(
            f'<defs>{RECT}<filter id="f"><feGaussianBlur stdDeviation="1"/></filter></defs>'
            '<use href="#r" filter="url(#f)"/>'
        ),
    )
    [made] = out.shapes()
    assert [c.tag.rpartition("}")[2] for c in made.find(f"{P}grpSpPr")] == ["xfrm", "effectLst"]


# ---------------------------------------------------------------------------------------------
# <symbol>
# ---------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("symbol", "use", "expected"),
    [
        ('viewBox="0 0 10 10"', 'width="50" height="50"', (u(100), u(100), u(50), u(50))),
        ('viewBox="0 0 10 10"', "", (u(100), u(100), u(10), u(10))),
        ('viewBox="5 5 10 10"', 'width="50" height="50"', (u(75), u(75), u(50), u(50))),
        ('viewBox="0,0,10,10"', 'width="20" height="20"', (u(100), u(100), u(20), u(20))),
        ("", 'width="50" height="50"', (u(100), u(100), u(10), u(10))),
        ('viewBox="0 0 10"', 'width="50" height="50"', (u(100), u(100), u(10), u(10))),
        ('viewBox="0 0 0 10"', 'width="50" height="50"', (u(100), u(100), u(10), u(10))),
        ('viewBox="a b c d"', 'width="50" height="50"', (u(100), u(100), u(10), u(10))),
    ],
)
def test_a_symbols_view_box_is_fitted_into_its_uses_box(
    tmp_path: Path, symbol: str, use: str, expected: tuple[int, int, int, int]
) -> None:
    out = lowered(
        tmp_path,
        slide(
            f'<defs><symbol id="s" {symbol}><rect width="10" height="10"/></symbol></defs>'
            f'<use href="#s" x="100" y="100" {use}/>'
        ),
    )
    assert box(first_sp(out.tree())) == expected


@pytest.mark.parametrize(
    ("ratio", "expected"),
    [
        ("", (u(112.5), u(100), u(25), u(50))),
        ('preserveAspectRatio="xMinYMin meet"', (u(100), u(100), u(25), u(50))),
        ('preserveAspectRatio="none"', (u(100), u(100), u(50), u(50))),
        ('preserveAspectRatio="xMidYMid slice"', (u(100), u(75), u(50), u(100))),
    ],
)
def test_a_tall_symbol_in_a_square_use_keeps_its_aspect_unless_told(
    tmp_path: Path, ratio: str, expected: tuple[int, int, int, int]
) -> None:
    out = lowered(
        tmp_path,
        slide(
            f'<defs><symbol id="s" viewBox="0 0 10 20" {ratio}><rect width="10" height="20"/>'
            '</symbol></defs><use href="#s" x="100" y="100" width="50" height="50"/>'
        ),
    )
    assert box(first_sp(out.tree())) == expected


def test_a_symbols_content_is_its_clone(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide(
            '<defs><symbol id="s"><rect width="1" height="1"/><circle r="1"/></symbol></defs>'
            '<use href="#s"/>'
        ),
    )
    assert names(out.tree()) == ["use 2", "rect 3", "circle 4"]


# ---------------------------------------------------------------------------------------------
# refusals
# ---------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("body", "said"),
    [
        ('<use href="#nowhere"/>', "#nowhere names no id in this slide"),
        (
            '<defs><linearGradient id="g"/></defs><use href="#g"/>',
            "names a <linearGradient>; it must name a shape, group, text, picture or symbol",
        ),
        ('<defs><use id="a" href="#a"/></defs><use href="#a"/>', "#a reaches itself"),
        (
            '<defs><use id="a" href="#b"/><use id="b" href="#a"/></defs><use href="#a"/>',
            "reaches itself",
        ),
    ],
)
def test_a_use_naming_nothing_or_itself_is_refused(tmp_path: Path, body: str, said: str) -> None:
    with pytest.raises(TargetError, match=said) as caught:
        lowered(tmp_path, slide(body))
    assert caught.value.code() == "OW_TARGET_REFERENCE_UNRESOLVED"


# ---------------------------------------------------------------------------------------------
# links
# ---------------------------------------------------------------------------------------------


def _rels(out: Any, n: int = 1) -> dict[str, dict[str, str]]:
    found = etree.fromstring(out.members[f"ppt/slides/_rels/slide{n}.xml.rels"])
    return {r.get("Id"): dict(r.attrib) for r in found}


def test_a_page_link_is_an_external_hyperlink_on_its_shape(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide(
            '<a href="https:example.org/x"><rect width="1" height="1"/></a>'
            '<rect width="1" height="1"/>'
        ),
    )
    assert clicks(out.tree()) == [("rect 2", ""), ("rect 3", None)]
    click = out.tree().find(f".//{A}hlinkClick")
    rel = _rels(out)[click.get(qn("r:id"))]
    assert (rel["Target"], rel["TargetMode"], rel["Type"].rpartition("/")[2]) == (
        "https:example.org/x",
        "External",
        "hyperlink",
    )
    [linked, _] = out.deck().slides[0].shapes
    assert linked.click_action.hyperlink.address == "https:example.org/x"


def test_a_unit_link_jumps_to_its_slide(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide('<a href="#u02"><rect width="1" height="1"/></a>'),
        slide('<a href="#u01"><rect width="1" height="1"/></a>'),
    )
    assert clicks(out.tree(1)) == [("rect 2", "ppaction://hlinksldjump")]
    click = out.tree(1).find(f".//{A}hlinkClick")
    rel = _rels(out)[click.get(qn("r:id"))]
    assert (rel["Target"], rel["Type"].rpartition("/")[2], "TargetMode" in rel) == (
        "slide2.xml",
        "slide",
        False,
    )
    first, second = out.deck().slides
    assert first.shapes[0].click_action.target_slide == second
    assert second.shapes[0].click_action.target_slide == first


def test_a_link_to_a_unit_the_deck_lacks_is_refused(tmp_path: Path) -> None:
    with pytest.raises(TargetError, match="#u09 names no slide in this deck") as caught:
        lowered(tmp_path, slide('<a href="#u09"><rect width="1" height="1"/></a>'))
    assert caught.value.code() == "OW_TARGET_REFERENCE_UNRESOLVED"


def test_every_shape_a_link_holds_clicks_and_an_inner_link_wins(tmp_path: Path) -> None:
    blob = png(2, 2)
    out = lowered(
        tmp_path,
        slide(
            '<a href="https:a.example/"><g><rect width="1" height="1"/>'
            f'<text y="20">t</text><image href="{picture_slot(blob)}" width="2" height="2"/>'
            '<a href="#u01"><circle r="1"/></a></g></a>'
        ),
        assets=[blob],
    )
    assert clicks(out.tree()) == [
        ("g 2", ""),
        ("rect 3", ""),
        ("text 4", ""),
        ("image 5", ""),
        ("circle 6", "ppaction://hlinksldjump"),
    ]


def test_two_links_to_one_page_are_one_relationship(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide(
            '<a href="https:example.org/x"><rect width="1" height="1"/></a>'
            '<a href="https:example.org/x"><rect width="1" height="1"/></a>'
        ),
    )
    ids = {c.get(qn("r:id")) for c in out.tree().iter(f"{A}hlinkClick")}
    assert len(ids) == 1
    assert [r["Target"] for r in _rels(out).values()].count("https:example.org/x") == 1
