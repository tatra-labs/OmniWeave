"""`compile.pptx.native`'s text: a `<text>` as one text box, a positioned line a paragraph. D797.

Each text is lowered through `lowered()` (tests/lowerkit.py) and read back from the package and
with python-pptx. The box's numbers follow `lower.text`'s model: a line is `LEADING` em for the
first and its baseline's distance after, the top is `ASCENT` of the first line above its
baseline, and the frame is `NOMINAL_ADVANCE` em a character of the widest line.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from lowerkit import (  # type: ignore[import-not-found]  # tests/lowerkit.py
    A,
    box,
    lowered,
    outline,
    slide,
    turn,
    u,
)
from lxml import etree  # type: ignore[attr-defined] -- lxml ships no stubs
from omniweave_core.errors import TargetError
from omniweave_target_pptx.driver import PptxCompiler
from omniweave_target_pptx.lower.text import run
from pptx.oxml.ns import qn

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def paragraphs(shape: Any) -> list[Any]:
    return shape.findall(f"{P}txBody/{A}p")


def runs(shape: Any) -> list[dict[str, str]]:
    """Each run's `a:rPr` attributes, in order."""
    return [dict(r.find(f"{A}rPr").attrib) for r in shape.iter(f"{A}r")]


def spacings(shape: Any) -> list[str]:
    return [s.get("val") for s in shape.iter(f"{A}spcPts")]


def text_of(shape: Any) -> list[str]:
    return ["".join(t.text or "" for t in p.iter(f"{A}t")) for p in paragraphs(shape)]


# ---------------------------------------------------------------------------------------------
# the box
# ---------------------------------------------------------------------------------------------


def test_a_text_is_a_text_box_in_drawingmls_order(tmp_path: Path) -> None:
    out = lowered(
        tmp_path, slide('<text data-ow-el="u01/e2" x="100" y="200" font-size="20">Hi</text>')
    )
    [shape] = out.shapes()
    assert outline(shape) == [
        "sp",
        "nvSpPr",
        "cNvPr id=2 name=u01/e2",
        "cNvSpPr txBox=1",
        "nvPr",
        "spPr",
        "xfrm",
        f"off x={u(100)} y={u(180.8)}",
        f"ext cx={u(24)} cy={u(24)}",
        "prstGeom prst=rect",
        "avLst",
        "noFill",
        "txBody",
        "bodyPr anchor=t bIns=0 lIns=0 rIns=0 tIns=0 wrap=none",
        "noAutofit",
        "lstStyle",
        "p",
        "pPr algn=l",
        "lnSpc",
        "spcPts val=1800",
        "r",
        "rPr lang=en-GB sz=1500",
        "solidFill",
        "srgbClr val=000000",
        "t",
        "endParaRPr lang=en-GB sz=1500",
    ]
    assert text_of(shape) == ["Hi"]


@pytest.mark.parametrize(
    ("body", "align", "left"),
    [
        ('<text x="100" y="40" font-size="20">Hi</text>', "l", u(100)),
        ('<text x="100" y="40" font-size="20" text-anchor="start">Hi</text>', "l", u(100)),
        ('<text x="100" y="40" font-size="20" text-anchor="middle">Hi</text>', "ctr", u(88)),
        ('<text x="100" y="40" font-size="20" text-anchor="end">Hi</text>', "r", u(76)),
        ('<g text-anchor="end"><text x="100" y="40" font-size="20">Hi</text></g>', "r", u(76)),
        ('<text y="40" font-size="20">Hi</text>', "l", 0),
        # a word the gate refuses (`pptx.enum@1`) is the initial `start`, never a guess
        ('<text x="100" y="40" font-size="20" text-anchor="left">Hi</text>', "l", u(100)),
    ],
)
def test_the_anchor_is_the_alignment_and_the_frames_edge_or_centre_at_x(
    tmp_path: Path, body: str, align: str, left: int
) -> None:
    shape = lowered(tmp_path, slide(body)).first()
    assert box(shape)[0] == left
    assert {p.find(f"{A}pPr").get("algn") for p in paragraphs(shape)} == {align}


@pytest.mark.parametrize(
    ("body", "width"),
    [
        ('<text y="40" font-size="20">a<tspan x="0" dy="30">bbbb</tspan></text>', u(48)),
        ('<text y="40" font-size="20" letter-spacing="2">a'
         '<tspan x="0" dy="30">bbbb</tspan></text>', u(56)),
        ('<text y="40" font-size="20">a<tspan font-size="10">bb</tspan></text>', u(24)),
        ('<text y="40"></text>', 0),
    ],
)  # fmt: skip
def test_the_frame_is_the_widest_lines_nominal_advance(
    tmp_path: Path, body: str, width: int
) -> None:
    assert box(lowered(tmp_path, slide(body)).first())[2] == width


# ---------------------------------------------------------------------------------------------
# lines
# ---------------------------------------------------------------------------------------------


def test_each_positioned_line_is_a_paragraph_at_its_baselines_distance(tmp_path: Path) -> None:
    shape = lowered(
        tmp_path,
        slide(
            '<text x="10" y="50" font-size="20">a<tspan x="10" dy="30">b</tspan>'
            '<tspan x="10" y="120">c</tspan><tspan dy="10">d</tspan></text>'
        ),
    ).first()
    assert text_of(shape) == ["a", "b", "c", "d"]
    assert spacings(shape) == ["1800", "2250", "3000", "750"]
    assert box(shape)[1::2] == (u(30.8), u(104))


def test_a_texts_own_dy_moves_its_first_baseline(tmp_path: Path) -> None:
    shape = lowered(tmp_path, slide('<text y="20" dy="5" font-size="20">a</text>')).first()
    assert box(shape)[1] == u(5.8)


def test_the_first_line_is_spaced_by_its_largest_run(tmp_path: Path) -> None:
    shape = lowered(
        tmp_path, slide('<text y="100" font-size="20">a<tspan font-size="40">B</tspan></text>')
    ).first()
    assert spacings(shape) == ["3600"]
    assert [r["sz"] for r in runs(shape)] == ["1500", "3000"]
    assert shape.find(f".//{A}endParaRPr").get("sz") == "3000"


def test_an_empty_line_keeps_its_size_at_its_end(tmp_path: Path) -> None:
    shape = lowered(
        tmp_path, slide('<text y="40" font-size="20">a<tspan x="0" dy="24" font-size="8"/></text>')
    ).first()
    empty = paragraphs(shape)[1]
    assert [etree.QName(c).localname for c in empty] == ["pPr", "endParaRPr"]
    assert empty.find(f"{A}endParaRPr").get("sz") == "600"


@pytest.mark.parametrize(
    ("body", "lines"),
    [
        ("<text y='10'>  a \n  b  </text>", ["a b"]),
        ("<text y='10' xml:space='preserve'>  a \n  b  </text>", ["  a    b  "]),
        ("<text y='10'>a <tspan>b</tspan> c</text>", ["a b c"]),
    ],
)
def test_whitespace_collapses_as_the_checker_reads_it(
    tmp_path: Path, body: str, lines: list[str]
) -> None:
    assert text_of(lowered(tmp_path, slide(body)).first()) == lines


# ---------------------------------------------------------------------------------------------
# runs
# ---------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("attrs", "expected"),
    [
        ("", {}),
        ('font-weight="bold"', {"b": "1"}),
        ('font-weight="600"', {"b": "1"}),
        ('font-weight="500"', {}),
        ('font-style="italic"', {"i": "1"}),
        ('text-decoration="underline"', {"u": "sng"}),
        ('text-decoration="line-through"', {"strike": "sngStrike"}),
        ('text-decoration="underline line-through"', {"u": "sng", "strike": "sngStrike"}),
        ('letter-spacing="2"', {"spc": "150"}),
        ('baseline-shift="super"', {"baseline": "30000"}),
        ('baseline-shift="sub"', {"baseline": "-25000"}),
        ('font-size="32"', {"sz": "2400"}),
        ('style="font-size: 32"', {"sz": "2400"}),
    ],
)
def test_a_runs_properties_are_drawingmls(
    tmp_path: Path, attrs: str, expected: dict[str, str]
) -> None:
    shape = lowered(tmp_path, slide(f'<text y="40">a<tspan {attrs}>b</tspan></text>')).first()
    plain, styled = runs(shape)
    assert plain == {"lang": "en-GB", "sz": "1200"}
    assert styled == {"lang": "en-GB", "sz": "1200", **expected}


def test_a_runs_properties_inherit_down_to_it(tmp_path: Path) -> None:
    shape = lowered(
        tmp_path,
        slide(
            '<g font-size="30" font-weight="bold"><text y="40" text-decoration="underline">'
            '<tspan font-style="italic">b</tspan></text></g>'
        ),
    ).first()
    assert runs(shape) == [{"b": "1", "i": "1", "lang": "en-GB", "sz": "2250", "u": "sng"}]


@pytest.mark.parametrize(
    ("stack", "face"),
    [
        ("Probe Sans, sans-serif", "Probe Sans"),
        ("'Other Serif', Arial", "Other Serif"),
        ('"Quoted Face"', "Quoted Face"),
        ("serif, Georgia", "Georgia"),
        ("sans-serif", None),
        ("SANS-SERIF, system-ui", None),
    ],
)
def test_the_face_is_the_first_family_that_is_not_generic(
    tmp_path: Path, stack: str, face: str | None
) -> None:
    quoted = stack.replace('"', "&quot;")
    shape = lowered(tmp_path, slide(f'<text y="40" font-family="{quoted}">a</text>')).first()
    properties = shape.find(f".//{A}rPr")
    faces = [(etree.QName(c).localname, c.get("typeface")) for c in properties][1:]
    expected = [] if face is None else [("latin", face), ("ea", face), ("cs", face)]
    assert faces == expected


def test_a_runs_fill_and_line_are_the_shapes_paint(tmp_path: Path) -> None:
    shape = lowered(
        tmp_path,
        slide(
            '<text y="40" fill="#FF0000" fill-opacity="0.5" stroke="#0000FF" stroke-width="2">'
            'a<tspan fill="none">b</tspan></text>'
        ),
    ).first()
    first, second = (r.find(f"{A}rPr") for r in shape.iter(f"{A}r"))
    assert outline(first) == [
        "rPr lang=en-GB sz=1200",
        f"ln w={u(2)}",
        "solidFill",
        "srgbClr val=0000FF",
        "solidFill",
        "srgbClr val=FF0000",
        "alpha val=50000",
    ]
    assert outline(second)[-1] == "noFill"


def test_a_run_with_no_language_names_none() -> None:
    doc = PptxCompiler().carrier("il/u01.svg", slide('<text y="40">a</text>'))
    [text] = doc.root().children()
    paragraph = etree.Element(qn("a:p"))
    run(paragraph, (doc.root(), text), text, "a", lang="", scale=1.0, definitions={})
    assert dict(paragraph.find(f"{A}r/{A}rPr").attrib) == {"sz": "1200"}


# ---------------------------------------------------------------------------------------------
# a matrix
# ---------------------------------------------------------------------------------------------


def test_a_scale_grows_the_box_its_sizes_and_its_lines(tmp_path: Path) -> None:
    shape = lowered(
        tmp_path,
        slide(
            '<text y="40" font-size="20" letter-spacing="2" stroke="#000000" '
            'transform="scale(2)">Hi</text>'
        ),
    ).first()
    assert box(shape) == (0, u(41.6), u(56), u(48))
    assert (runs(shape)[0]["sz"], runs(shape)[0]["spc"]) == ("3000", "300")
    assert spacings(shape) == ["3600"]
    assert shape.find(f".//{A}endParaRPr").get("sz") == "3000"
    assert shape.find(f".//{A}rPr/{A}ln").get("w") == str(u(2))


def test_a_non_scaling_runs_line_keeps_its_width(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide(
            '<text y="40" font-size="20" stroke="#000000" vector-effect="non-scaling-stroke" '
            'transform="scale(2)">Hi</text>'
        ),
    )
    assert out.first().find(f".//{A}rPr/{A}ln").get("w") == str(u(1))


def test_a_turned_text_box_turns_about_its_centre(tmp_path: Path) -> None:
    shape = lowered(
        tmp_path, slide('<text y="40" font-size="20" transform="rotate(90)">Hi</text>')
    ).first()
    assert turn(shape) == ("5400000", None, None)
    assert box(shape) == (u(-44.8), 0, u(24), u(24))


# ---------------------------------------------------------------------------------------------
# refusals, and python-pptx reading it back
# ---------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("body", "said"),
    [
        ('<text y="40">a<tspan dx="4">b</tspan></text>', "dx='4' moves a run"),
        ('<text y="40" dx="-1">a</text>', "dx='-1' moves a run"),
        ('<text y="40">a<tspan x="0" dy="-5">b</tspan></text>', "is not below the line before"),
        ('<text y="40">a<tspan x="0" y="40">b</tspan></text>', "is not below the line before"),
    ],
)
def test_a_place_one_text_box_cannot_draw_is_refused(tmp_path: Path, body: str, said: str) -> None:
    with pytest.raises(TargetError, match=said) as caught:
        lowered(tmp_path, slide(body))
    assert caught.value.code() == "OW_TARGET_TEXT_INVALID"


def test_python_pptx_reads_the_text_back(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide(
            '<text data-ow-el="u01/e2" x="64" y="120" font-family="Probe Sans" font-size="40">'
            'Revenue <tspan font-weight="bold">rose</tspan><tspan x="64" dy="48">12%</tspan>'
            "</text>"
        ),
    )
    [page] = out.deck().slides
    [shape] = page.shapes
    assert (shape.name, shape.has_text_frame) == ("u01/e2", True)
    first, second = shape.text_frame.paragraphs
    assert (first.text, second.text) == ("Revenue rose", "12%")
    plain, bold = first.runs
    assert (plain.font.bold, bold.font.bold, bold.font.size.pt, bold.font.name) == (
        None,
        True,
        30.0,
        "Probe Sans",
    )
