"""`compile.pptx.native`'s freeforms: paths, polygons and polylines as `a:custGeom`. D796.

`absolute()` and `arc()` are read directly, the way SVG 1.1 section 8.3 and its implementation
notes (F.6.5, F.6.6) draw a path, and each freeform through `lowered()` (tests/lowerkit.py), read
back from the package. A point is in EMU from the freeform's box, and the box holds every point,
control points included.
"""

from __future__ import annotations

import math
from pathlib import Path
from typing import Any

import pytest
from lowerkit import (  # type: ignore[import-not-found]  # tests/lowerkit.py
    INK,
    A,
    P,
    box,
    line_of,
    lowered,
    outline,
    slide,
    turn,
    u,
)
from lxml import etree  # type: ignore[attr-defined] -- lxml ships no stubs
from omniweave_core.errors import TargetError
from omniweave_target_pptx.lower.paths import absolute, arc
from omniweave_target_pptx.rules.path_data import segments
from pptx.enum.shapes import MSO_SHAPE_TYPE

K = 4 / 3 * math.tan(math.pi / 8)
"""A quarter turn's cubic: its controls sit `K` radii along the tangents."""
ROOT2 = math.sqrt(2)


def steps(shape: Any) -> list[tuple[str, list[tuple[int, int]]]]:
    """A freeform's path: each step's name and its points."""
    path = shape.find(f"{P}spPr/{A}custGeom/{A}pathLst/{A}path")
    return [
        (etree.QName(s).localname, [(int(p.get("x")), int(p.get("y"))) for p in s.iter(f"{A}pt")])
        for s in path
    ]


def drawn(d: str) -> list[tuple[str, tuple[tuple[float, float], ...]]]:
    return absolute(segments(d))


# ---------------------------------------------------------------------------------------------
# polygons and polylines
# ---------------------------------------------------------------------------------------------


def test_a_polygon_is_one_closed_freeform_in_drawingmls_order(tmp_path: Path) -> None:
    out = lowered(
        tmp_path,
        slide('<polygon data-ow-el="u01/e2" points="10,10 110,10 60,60" fill="#1F4E79"/>'),
    )
    [shape] = out.shapes()
    assert outline(shape) == [
        "sp",
        "nvSpPr",
        "cNvPr id=2 name=u01/e2",
        "cNvSpPr",
        "nvPr",
        "spPr",
        "xfrm",
        f"off x={u(10)} y={u(10)}",
        f"ext cx={u(100)} cy={u(50)}",
        "custGeom",
        "avLst",
        "gdLst",
        "ahLst",
        "cxnLst",
        "rect b=b l=l r=r t=t",
        "pathLst",
        f"path h={u(50)} w={u(100)}",
        "moveTo",
        "pt x=0 y=0",
        "lnTo",
        f"pt x={u(100)} y=0",
        "lnTo",
        f"pt x={u(50)} y={u(50)}",
        "close",
        "solidFill",
        "srgbClr val=1F4E79",
        "ln cap=flat w=0",
        "noFill",
        "miter lim=800000",
    ]


def test_a_polyline_is_open_and_still_filled(tmp_path: Path) -> None:
    shape = lowered(tmp_path, slide('<polyline points="0,0 10 10,20,0"/>')).first()
    assert steps(shape) == [
        ("moveTo", [(0, 0)]),
        ("lnTo", [(u(10), u(10))]),
        ("lnTo", [(u(20), 0)]),
    ]
    assert outline(shape.find(f"{P}spPr")[2]) == ["solidFill", "srgbClr val=000000"]


@pytest.mark.parametrize(
    ("transform", "expected", "points"),
    [
        (
            "rotate(90)",
            (u(-10), 0, u(10), u(10)),
            [(u(10), 0), (u(10), u(10)), (0, u(10))],
        ),
        (
            "scale(-1 1)",
            (u(-10), 0, u(10), u(10)),
            [(u(10), 0), (0, 0), (u(10), u(10))],
        ),
        (
            "translate(5 6) scale(2)",
            (u(5), u(6), u(20), u(20)),
            [(0, 0), (u(20), 0), (u(20), u(20))],
        ),
    ],
)
def test_a_freeforms_points_take_its_matrix_and_it_keeps_no_turn(
    tmp_path: Path,
    transform: str,
    expected: tuple[int, int, int, int],
    points: list[tuple[int, int]],
) -> None:
    body = (
        f'<polygon points="0,0 10,0 10,10" transform="{transform}"/>'
        if transform != "scale(-1 1)"
        else f'<polygon points="0,0 10,0 0,10" transform="{transform}"/>'
    )
    shape = lowered(tmp_path, slide(body)).first()
    assert box(shape) == expected
    assert turn(shape) == (None, None, None)
    assert [p for _, found in steps(shape) for p in found] == points


@pytest.mark.parametrize("points", ["0,0 10", "", "0,0 1e3,2"])
def test_points_that_are_not_pairs_are_refused(tmp_path: Path, points: str) -> None:
    with pytest.raises(TargetError, match=r"<polygon> at .*: points are not x,y pairs") as caught:
        lowered(tmp_path, slide(f'<polygon points="{points}"/>'))
    assert caught.value.code() == "OW_TARGET_GEOMETRY_INVALID"


# ---------------------------------------------------------------------------------------------
# path data, as SVG draws it
# ---------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("d", "ops"),
    [
        ("M 1 2 L 3 4", [("M", ((1, 2),)), ("L", ((3, 4),))]),
        ("M 1 2 3 4 5 6", [("M", ((1, 2),)), ("L", ((3, 4),)), ("L", ((5, 6),))]),
        ("m 1 2 3 4", [("M", ((1, 2),)), ("L", ((4, 6),))]),
        (
            "M 1 1 l 2 3 h 4 v 5",
            [("M", ((1, 1),)), ("L", ((3, 4),)), ("L", ((7, 4),)), ("L", ((7, 9),))],
        ),
        ("M 1 1 H 4 V 5", [("M", ((1, 1),)), ("L", ((4, 1),)), ("L", ((4, 5),))]),
        (
            "M 1 1 c 1 1 2 2 3 3 1 1 2 2 3 3",
            [("M", ((1, 1),)), ("C", ((2, 2), (3, 3), (4, 4))), ("C", ((5, 5), (6, 6), (7, 7)))],
        ),
        (
            "M 0 0 C 1 2 3 4 5 6 S 9 9 10 10",
            [("M", ((0, 0),)), ("C", ((1, 2), (3, 4), (5, 6))), ("C", ((7, 8), (9, 9), (10, 10)))],
        ),
        (
            "M 0 0 c 1 2 3 4 5 6 s 4 3 5 4",
            [("M", ((0, 0),)), ("C", ((1, 2), (3, 4), (5, 6))), ("C", ((7, 8), (9, 9), (10, 10)))],
        ),
        (
            "M 0 0 L 5 5 S 9 9 10 10",
            [("M", ((0, 0),)), ("L", ((5, 5),)), ("C", ((5, 5), (9, 9), (10, 10)))],
        ),
        (
            "M 0 0 S 2 2 3 3 S 5 5 6 6",
            [("M", ((0, 0),)), ("C", ((0, 0), (2, 2), (3, 3))), ("C", ((4, 4), (5, 5), (6, 6)))],
        ),
        (
            "M 0 0 Q 1 1 2 2 S 4 4 5 5",
            [("M", ((0, 0),)), ("Q", ((1, 1), (2, 2))), ("C", ((2, 2), (4, 4), (5, 5)))],
        ),
        (
            "M 0 0 Q 10 0 10 10 T 20 20",
            [("M", ((0, 0),)), ("Q", ((10, 0), (10, 10))), ("Q", ((10, 20), (20, 20)))],
        ),
        (
            "M 0 0 q 10 0 10 10 t 10 10 10 10",
            [
                ("M", ((0, 0),)),
                ("Q", ((10, 0), (10, 10))),
                ("Q", ((10, 20), (20, 20))),
                ("Q", ((30, 20), (30, 30))),
            ],
        ),
        (
            "M 0 0 C 1 1 2 2 3 3 T 5 5",
            [("M", ((0, 0),)), ("C", ((1, 1), (2, 2), (3, 3))), ("Q", ((3, 3), (5, 5)))],
        ),
        (
            "M 1 1 L 5 1 Z L 9 9",
            [("M", ((1, 1),)), ("L", ((5, 1),)), ("Z", ()), ("M", ((1, 1),)), ("L", ((9, 9),))],
        ),
        (
            "M 1 1 L 5 1 z m 2 2 l 1 0",
            [("M", ((1, 1),)), ("L", ((5, 1),)), ("Z", ()), ("M", ((3, 3),)), ("L", ((4, 3),))],
        ),
        (
            "M 0 0 C 1 1 2 2 3 3 Z S 5 5 6 6",
            [
                ("M", ((0, 0),)),
                ("C", ((1, 1), (2, 2), (3, 3))),
                ("Z", ()),
                ("M", ((0, 0),)),
                ("C", ((0, 0), (5, 5), (6, 6))),
            ],
        ),
        ("M 0 0 Z Z", [("M", ((0, 0),)), ("Z", ()), ("Z", ())]),
        ("M 3 3 A 5 5 0 0 1 3 3", [("M", ((3, 3),))]),
        ("M 0 0 A 0 5 0 0 1 10 0", [("M", ((0, 0),)), ("L", ((10, 0),))]),
        ("M 1 1 a 5 0 0 0 1 10 0", [("M", ((1, 1),)), ("L", ((11, 1),))]),
    ],
)
def test_path_data_is_drawn_as_svg_draws_it(
    d: str, ops: list[tuple[str, tuple[tuple[float, float], ...]]]
) -> None:
    assert drawn(d) == ops


def test_a_smooth_curve_after_an_arc_starts_at_the_current_point() -> None:
    *_, last = drawn("M 0 0 C 1 1 2 2 3 3 A 1 1 0 0 1 4 3 S 5 5 6 6")
    assert last == ("C", ((4.0, 3.0), (5.0, 5.0), (6.0, 6.0)))


def test_an_arc_is_cubics_that_end_where_it_ends() -> None:
    ops = drawn("M 0 0 a 10 10 0 0 1 20 0 A 10 10 0 1 1 20 0.5")
    assert [kind for kind, _ in ops] == ["M", "C", "C", "C", "C", "C", "C"]
    assert ops[2][1][-1] == (20.0, 0.0)
    assert ops[-1][1][-1] == (20.0, 0.5)


# ---------------------------------------------------------------------------------------------
# arcs (SVG 1.1 F.6.5 and F.6.6)
# ---------------------------------------------------------------------------------------------


def _close(found: list[Any], expected: list[Any]) -> None:
    flat = [v for cubic in found for p in cubic for v in p]
    want = [v for cubic in expected for p in cubic for v in p]
    assert flat == pytest.approx(want, abs=1e-9)


def test_a_quarter_circle_is_one_cubic_with_its_controls_on_the_tangents() -> None:
    _close(
        arc((0, 0), (10, 10), 10, 10, 0, large=False, sweep=True),
        [((10 * K, 0), (10, 10 - 10 * K), (10, 10))],
    )


def test_the_sweep_flag_picks_the_other_centre() -> None:
    _close(
        arc((0, 0), (10, 10), 10, 10, 0, large=False, sweep=False),
        [((0, 10 * K), (10 - 10 * K, 10), (10, 10))],
    )


def test_the_large_flag_takes_three_quarters_in_three_pieces() -> None:
    """About (10, 0), the other centre: up through (10, -10) and round by (20, 0)."""
    found = arc((0, 0), (10, 10), 10, 10, 0, large=True, sweep=True)
    assert len(found) == 3
    _close(
        found,
        [
            ((0, -10 * K), (10 - 10 * K, -10), (10, -10)),
            ((10 + 10 * K, -10), (20, -10 * K), (20, 0)),
            ((20, 10 * K), (10 + 10 * K, 10), (10, 10)),
        ],
    )


@pytest.mark.parametrize(
    ("end", "large", "pieces"),
    [((20, 0), False, 2), ((0.001, 0), True, 4), ((0.001, 0), False, 1)],
)
def test_a_sweep_is_cut_into_quarter_turns_at_most(
    end: tuple[float, float], large: bool, pieces: int
) -> None:
    assert len(arc((0, 0), end, 10, 10, 0, large=large, sweep=True)) == pieces


def test_a_large_arc_against_the_sweep_takes_three_quarters_the_other_way() -> None:
    found = arc((0, 0), (10, 10), 10, 10, 0, large=True, sweep=False)
    ends = [v for cubic in found for v in cubic[-1]]
    assert ends == pytest.approx([-10, 10, 0, 20, 10, 10])


@pytest.mark.parametrize("degrees", [15, 20, 30])
def test_a_quarter_a_hair_over_by_rounding_is_still_one_piece(degrees: int) -> None:
    """A turned ellipse's quarter sweeps a quarter turn and a few ulps."""
    c, s = math.cos(math.radians(degrees)), math.sin(math.radians(degrees))
    assert len(arc((10 * c, 10 * s), (-5 * s, 5 * c), 10, 5, degrees, large=False, sweep=True)) == 1


def test_radii_too_small_to_reach_are_scaled_until_they_do() -> None:
    _close(
        arc((0, 0), (20, 0), 1, 1, 0, large=False, sweep=True),
        arc((0, 0), (20, 0), 10, 10, 0, large=False, sweep=True),
    )


def test_a_rotated_ellipses_quarter_runs_from_its_major_vertex_to_its_minor() -> None:
    """rx 10 and ry 5, turned 45 degrees about the origin: a quarter from (10, 0) to (0, 5)."""

    def turned(x: float, y: float) -> tuple[float, float]:
        return ((x - y) / ROOT2, (x + y) / ROOT2)

    _close(
        arc(turned(10, 0), turned(0, 5), 10, 5, 45, large=False, sweep=True),
        [(turned(10, 5 * K), turned(10 * K, 5), turned(0, 5))],
    )


# ---------------------------------------------------------------------------------------------
# paths, lowered
# ---------------------------------------------------------------------------------------------


def test_a_paths_box_holds_its_control_points(tmp_path: Path) -> None:
    out = lowered(tmp_path, slide('<path data-ow-el="u01/e2" d="M 10 10 C 20 0 30 0 40 10 Z"/>'))
    [shape] = out.shapes()
    assert box(shape) == (u(10), 0, u(30), u(10))
    assert steps(shape) == [
        ("moveTo", [(0, u(10))]),
        ("cubicBezTo", [(u(10), 0), (u(20), 0), (u(30), u(10))]),
        ("close", []),
    ]
    path = shape.find(f"{P}spPr/{A}custGeom/{A}pathLst/{A}path")
    assert (path.get("w"), path.get("h")) == (str(u(30)), str(u(10)))
    assert out.deck().slides[0].shapes[0].shape_type == MSO_SHAPE_TYPE.FREEFORM


def test_a_quadratic_is_drawingmls_own(tmp_path: Path) -> None:
    shape = lowered(tmp_path, slide('<path d="M 0 0 Q 10 10 20 0"/>')).first()
    assert steps(shape) == [("moveTo", [(0, 0)]), ("quadBezTo", [(u(10), u(10)), (u(20), 0)])]


@pytest.mark.parametrize(
    ("d", "expected", "across"),
    [
        ("M 0 5 L 10 5", (0, u(5), u(10), 0), (str(u(10)), "1")),
        ("M 5 0 L 5 10", (u(5), 0, 0, u(10)), ("1", str(u(10)))),
    ],
)
def test_an_axis_with_no_extent_is_a_path_one_emu_across(
    tmp_path: Path, d: str, expected: tuple[int, int, int, int], across: tuple[str, str]
) -> None:
    shape = lowered(tmp_path, slide(f'<path d="{d}" {INK}/>')).first()
    assert box(shape) == expected
    path = shape.find(f"{P}spPr/{A}custGeom/{A}pathLst/{A}path")
    assert (path.get("w"), path.get("h")) == across


def test_an_arc_in_a_path_is_its_cubics(tmp_path: Path) -> None:
    shape = lowered(tmp_path, slide('<path d="M 0 0 A 10 10 0 0 1 20 0"/>')).first()
    assert [name for name, _ in steps(shape)] == ["moveTo", "cubicBezTo", "cubicBezTo"]
    assert steps(shape)[-1][1][-1] == (u(20), u(10))


def test_a_paths_stroke_grows_with_its_scale(tmp_path: Path) -> None:
    shape = lowered(tmp_path, slide(f'<path d="M 0 0 L 10 0" {INK} transform="scale(2)"/>')).first()
    assert line_of(shape).get("w") == str(u(2))


@pytest.mark.parametrize(
    ("d", "said"),
    [("L 0 0", "path data starts with M or m"), ("", "the path data is empty")],
)
def test_path_data_the_gate_would_refuse_is_refused(tmp_path: Path, d: str, said: str) -> None:
    with pytest.raises(TargetError, match=f"<path> at .*: {said}") as caught:
        lowered(tmp_path, slide(f'<path d="{d}"/>'))
    assert caught.value.code() == "OW_TARGET_PATH_INVALID"
